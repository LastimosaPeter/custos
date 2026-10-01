import csv
import html
import io
import json
import math
import os
import random
import re
import secrets
import threading
import time
from datetime import datetime, timedelta
from functools import wraps

from dotenv import load_dotenv
from flask import (
    Flask, Response, abort, flash, g, jsonify, redirect, render_template,
    request, session, url_for, send_from_directory, send_file
)
from werkzeug.middleware.proxy_fix import ProxyFix
from werkzeug.security import check_password_hash

from notebook_workbench import load_workbench_record

from db import APP_TZ, DATABASE_ENGINE, connect, ensure_db_initialized, init_db, iso_now, unique_session_key
from item_analysis import build_item_analysis, build_programming_analysis
from google_integration import (
    STUDENT_GOOGLE_LOGIN_REQUIRED, csp_additions, google_profile_names, google_student_identity, roster_entry,
)

# Pages that render a Google sign-in button or the Classroom connector.
GOOGLE_SIGNIN_ENDPOINTS = {"student_login", "admin_login", "nextgen.custom_assessment", "admin_dashboard", "admin_batch", "nextgen.workspace"}

load_dotenv()

app = Flask(__name__)
_secret_key = os.getenv("SECRET_KEY", "").strip()
_admin_password = os.getenv("ADMIN_PASSWORD", "").strip()
if not _secret_key or _secret_key.startswith("replace-"):
    raise RuntimeError("Set a strong SECRET_KEY environment variable before starting Custos.")
if not _admin_password or _admin_password.startswith("replace-"):
    raise RuntimeError("Set a strong ADMIN_PASSWORD environment variable before starting Custos.")
app.secret_key = _secret_key
_hosted_https_default = "1" if (os.getenv("VERCEL") or os.getenv("RENDER") or os.getenv("RENDER_EXTERNAL_URL")) else "0"
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Strict",
    SESSION_COOKIE_SECURE=os.getenv("SESSION_COOKIE_SECURE", _hosted_https_default) == "1",
    PERMANENT_SESSION_LIFETIME=timedelta(days=7),
    MAX_CONTENT_LENGTH=32 * 1024 * 1024,
)
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1)

APP_NAME = os.getenv("APP_NAME", "Custos")
APP_VERSION = "1.6.0.h"
APP_RELEASE_SPECIES = "Hercules"
APP_RELEASE_SCIENTIFIC_NAME = "Dynastes hercules"
APP_RELEASE_COMMON_NAME = "Hercules beetle release"
APP_ASSET_REVISION = "1.6.0-h-hercules-workbench-font-r31"

ALLOWED_EMAIL_DOMAIN = os.getenv("ALLOWED_EMAIL_DOMAIN", "adnu.edu.ph").lower()
SUSPICIOUS_EVENTS = {
    "inactivity", "inactivity_no_response", "blocked_shortcut",
    "contextmenu", "beforeunload"
}
TEMP_LOCK_SECONDS = 15
MAX_SECURITY_VIOLATIONS = 3
SECURITY_DEDUPE_SECONDS = 3
ADMIN_LOGIN_MAX_FAILURES = 5
ADMIN_LOGIN_WINDOW_SECONDS = 300
ADMIN_LOGIN_BLOCK_SECONDS = 300
_admin_login_guard = {}
_admin_login_guard_lock = threading.Lock()


def _admin_login_guard_key():
    return str(request.remote_addr or "unknown")[:120]


def _admin_login_retry_after(key):
    now = time.monotonic()
    with _admin_login_guard_lock:
        state = _admin_login_guard.get(key)
        if not state:
            return 0
        blocked_until = float(state.get("blocked_until") or 0)
        if blocked_until > now:
            return max(1, int(blocked_until - now))
        attempts = [ts for ts in state.get("attempts", []) if now - ts <= ADMIN_LOGIN_WINDOW_SECONDS]
        if attempts:
            state["attempts"] = attempts
            state["blocked_until"] = 0
        else:
            _admin_login_guard.pop(key, None)
        return 0


def _record_admin_login_failure(key):
    now = time.monotonic()
    with _admin_login_guard_lock:
        state = _admin_login_guard.setdefault(key, {"attempts": [], "blocked_until": 0})
        state["attempts"] = [ts for ts in state.get("attempts", []) if now - ts <= ADMIN_LOGIN_WINDOW_SECONDS]
        state["attempts"].append(now)
        if len(state["attempts"]) >= ADMIN_LOGIN_MAX_FAILURES:
            state["blocked_until"] = now + ADMIN_LOGIN_BLOCK_SECONDS
            state["attempts"].clear()


def _clear_admin_login_failures(key):
    with _admin_login_guard_lock:
        _admin_login_guard.pop(key, None)

STUDENT_SECTIONS = {
    "ZT": {"11", "12", "13"},
    "ZS": {"11"},
}
CSEC303_SECTIONS = {
    "ZC": {"32"},
}

SUPPORTED_SUBJECTS = {
    "CSDC101": {"name": "Fundamentals of Programming", "logo": "img/csdc101-logo.webp"},
    "CSEC303": {"name": "Digital Image Processing", "logo": "img/csec303-logo.webp"},
}



def _workbench_asset_url_map(assessment_id, asset_rows):
    return {
        str(row["asset_path"]): url_for("assessment_workbench_asset", assessment_id=assessment_id, asset_path=row["asset_path"])
        for row in asset_rows
    }


def _render_notebook_markdown(source, asset_urls=None):
    """Render a small safe Markdown subset used by notebook workbenches.

    Relative Markdown images are rewritten to the assessment's database-backed
    asset endpoint, making uploaded notebook ZIPs portable on Vercel, Render,
    local SQLite, and PostgreSQL without writing to the server filesystem.
    """
    source = str(source or "")
    asset_urls = asset_urls or {}
    image_tokens = {}

    def stash_image(match):
        alt = html.escape(match.group(1) or "Notebook image", quote=True)
        raw = str(match.group(2) or "").strip().split("#", 1)[0]
        target = asset_urls.get(raw) or asset_urls.get(raw.lstrip("./"))
        if not target:
            return html.escape(match.group(0), quote=True)
        token = f"__CUSTOS_IMG_{len(image_tokens)}__"
        image_tokens[token] = f'<img src="{html.escape(target, quote=True)}" alt="{alt}" loading="lazy">'
        return token

    source = re.sub(r"!\[([^\]]*)\]\(([^)]+)\)", stash_image, source)
    lines = source.splitlines()
    out = []
    in_list = False

    def inline(value):
        value = html.escape(value, quote=True)
        value = re.sub(r"`([^`]+)`", r"<code>\1</code>", value)
        value = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", value)
        for token, image_html in image_tokens.items():
            value = value.replace(html.escape(token), image_html)
        return value

    for raw in lines:
        stripped = raw.rstrip().strip()
        if not stripped:
            if in_list:
                out.append("</ul>")
                in_list = False
            continue
        if stripped.startswith("### "):
            if in_list: out.append("</ul>"); in_list = False
            out.append(f"<h4>{inline(stripped[4:])}</h4>")
        elif stripped.startswith("## "):
            if in_list: out.append("</ul>"); in_list = False
            out.append(f"<h3>{inline(stripped[3:])}</h3>")
        elif stripped.startswith("# "):
            if in_list: out.append("</ul>"); in_list = False
            out.append(f"<h2>{inline(stripped[2:])}</h2>")
        elif stripped.startswith(("- ", "* ")):
            if not in_list: out.append("<ul>"); in_list = True
            out.append(f"<li>{inline(stripped[2:])}</li>")
        elif re.match(r"^\d+\.\s+", stripped):
            if not in_list: out.append("<ul>"); in_list = True
            out.append(f"<li>{inline(re.sub(r'^\d+\.\s+', '', stripped))}</li>")
        else:
            if in_list: out.append("</ul>"); in_list = False
            out.append(f"<p>{inline(stripped)}</p>")
    if in_list:
        out.append("</ul>")
    return "".join(out)


def _section_key(title, used):
    case = re.search(r"\bCASE\s*0?(\d+)\b", title.upper())
    if case:
        base = f"CASE{int(case.group(1)):02d}"
    elif "FINAL" in title.upper():
        base = "FINAL"
    else:
        base = re.sub(r"[^A-Z0-9]+", "_", title.upper()).strip("_")[:24] or "SECTION"
    key = base
    n = 2
    while key in used:
        key = f"{base}_{n}"; n += 1
    used.add(key)
    return key


def _parse_notebook_sections(notebook, asset_urls=None, metadata=None):
    metadata = metadata or {}
    previews = metadata.get("previews") if isinstance(metadata.get("previews"), dict) else {}
    setup_cells = {int(v) for v in metadata.get("setup_cells", []) if str(v).isdigit()} if isinstance(metadata.get("setup_cells"), list) else set()
    sections = []
    current = {"key": "GUIDE", "title": metadata.get("title") or "Notebook Guide", "tab": "Guide", "cells": [], "preview": None}
    sections.append(current)
    used = {"GUIDE"}
    first_heading_seen = False

    for cell_index, cell in enumerate(notebook.get("cells", [])):
        source = "".join(cell.get("source", []))
        if cell.get("cell_type") == "markdown":
            heading = re.search(r"(?m)^#{1,2}\s+(.+?)\s*$", source)
            if heading:
                title = re.sub(r"[*_`]+", "", heading.group(1)).strip()
                # The first H1 is normally the notebook title and remains in Guide.
                if first_heading_seen or re.search(r"\bCASE\s*0?\d+\b|FINAL\s+CHECK", title, re.I):
                    key = _section_key(title, used)
                    current = {"key": key, "title": title, "tab": (re.search(r"CASE\s*0?(\d+)", title, re.I).group(1).zfill(2) if re.search(r"CASE\s*0?(\d+)", title, re.I) else title[:16]), "cells": [], "preview": None}
                    sections.append(current)
                first_heading_seen = True
        is_setup = bool(cell_index in setup_cells or (cell.get("cell_type") == "code" and ("def gray" in source or "import numpy" in source) and len(current["cells"]) <= 2))
        current["cells"].append({
            "id": cell_index,
            "type": cell.get("cell_type", "markdown"),
            "source": source.rstrip(),
            "is_setup": is_setup,
            "html": _render_notebook_markdown(source, asset_urls) if cell.get("cell_type") == "markdown" else "",
        })

    for section in sections:
        preview_name = previews.get(section["key"]) or previews.get(section["title"])
        if preview_name and asset_urls:
            section["preview"] = asset_urls.get(str(preview_name)) or asset_urls.get(str(preview_name).lstrip("./"))
    return [s for s in sections if s["cells"]]


def _load_static_workbench(assessment):
    """Load a GitHub-deployable workbench from static/workbenches/<assessment-slug>/.

    This is the file-based companion to database uploads. It keeps PHANTOM-303
    and future public notebooks template-driven instead of hard-coding one lab.
    """
    slug = str(assessment["slug"] or "").strip()
    if not slug or not re.fullmatch(r"[a-z0-9][a-z0-9_-]{0,120}", slug):
        return None
    rel_root = f"workbenches/{slug}"
    base = os.path.join(app.static_folder, "workbenches", slug)
    descriptor_path = os.path.join(base, "custos-workbench.json")
    if not os.path.isfile(descriptor_path):
        return None
    try:
        with open(descriptor_path, "r", encoding="utf-8") as fh:
            descriptor = json.load(fh)
        if not isinstance(descriptor, dict):
            return None
        notebook_name = str(descriptor.get("notebook") or "").strip()
        if not notebook_name:
            choices = [name for name in os.listdir(base) if name.lower().endswith(".ipynb")]
            if len(choices) != 1:
                return None
            notebook_name = choices[0]
        if os.path.basename(notebook_name) != notebook_name:
            return None
        notebook_path = os.path.join(base, notebook_name)
        with open(notebook_path, "r", encoding="utf-8") as fh:
            notebook = json.load(fh)

        notebook_meta = notebook.get("metadata", {}).get("custos", {})
        metadata = dict(notebook_meta) if isinstance(notebook_meta, dict) else {}
        for key in ("title", "subtitle", "previews", "setup_cells"):
            if key in descriptor:
                metadata[key] = descriptor[key]

        asset_urls = {}
        assets = []
        for root, _dirs, names in os.walk(base):
            for name in names:
                abs_path = os.path.join(root, name)
                rel = os.path.relpath(abs_path, base).replace(os.sep, "/")
                if rel == "custos-workbench.json" or rel == notebook_name or rel.lower().endswith((".zip", ".md", ".txt", ".ipynb")):
                    continue
                url = url_for("static", filename=f"{rel_root}/{rel}", v=APP_ASSET_REVISION)
                asset_urls[rel] = url
                assets.append({"path": rel, "url": url})

        package_name = str(descriptor.get("package") or "").strip()
        package_path = os.path.join(base, package_name) if package_name and os.path.basename(package_name) == package_name else ""
        package_url = url_for("static", filename=f"{rel_root}/{package_name}") if package_path and os.path.isfile(package_path) else None
        return {
            "sections": _parse_notebook_sections(notebook, asset_urls, metadata),
            "title": metadata.get("title") or os.path.splitext(notebook_name)[0],
            "subtitle": metadata.get("subtitle") or "Scientific Python Workbench",
            "assets": assets,
            "notebook_url": url_for("static", filename=f"{rel_root}/{notebook_name}"),
            "package_url": package_url,
            "filename": notebook_name,
            "source": "static",
        }
    except (OSError, ValueError, json.JSONDecodeError):
        return None


def get_assessment_workbench(conn, assessment):
    """Load the database workbench first, then a GitHub/static workbench template."""
    if not assessment:
        return None
    record = load_workbench_record(conn, int(assessment["id"]))
    if record:
        nb_row, assets = record
        try:
            notebook = json.loads(nb_row["notebook_json"])
            metadata = json.loads(nb_row["metadata_json"] or "{}")
        except Exception:
            return None
        asset_urls = _workbench_asset_url_map(assessment["id"], assets)
        sections = _parse_notebook_sections(notebook, asset_urls, metadata)
        return {
            "sections": sections,
            "title": metadata.get("title") or os.path.splitext(nb_row["filename"])[0],
            "subtitle": metadata.get("subtitle") or "Scientific Python Workbench",
            "assets": [{"path": row["asset_path"], "url": asset_urls[row["asset_path"]]} for row in assets],
            "notebook_url": url_for("assessment_workbench_notebook", assessment_id=assessment["id"]),
            "package_url": url_for("assessment_workbench_package", assessment_id=assessment["id"]),
            "filename": nb_row["filename"],
            "source": "database",
        }
    return _load_static_workbench(assessment)

def normalize_subject_code(value):
    code = str(value or "").strip().upper()
    return code if code in SUPPORTED_SUBJECTS else "CSDC101"

THEME_COOKIE_NAME = "custos_theme"
THEME_COOKIE_MAX_AGE = 60 * 60 * 24 * 365


def get_saved_theme():
    """Theme preference is intentionally client/session scoped.

    Older Custos builds stored this tiny UI preference in PostgreSQL, which meant
    every rendered page paid for an unnecessary database round trip. Keeping the
    preference in Flask's signed session makes page rendering faster and works on
    Render, Vercel, Docker, and local SQLite without changing assessment data.
    """
    theme = session.get("ui_theme", "light")
    return theme if theme in {"light", "dark"} else "light"


# Avoid a full schema migration on every Gunicorn/Vercel cold start. New or
# upgraded databases are still initialized automatically; established databases
# only perform a tiny schema-version probe.
if os.getenv("AUTO_INIT_DB", "1") == "1":
    ensure_db_initialized(
        admin_username=os.getenv("ADMIN_USERNAME", "admin"),
        admin_password=_admin_password,
        force=os.getenv("FORCE_DB_INIT", "0") == "1",
    )


def normalize_short_answer(value):
    """Case-insensitive short-answer normalization.

    Spaces, dashes, punctuation, and other non-alphanumeric characters are
    ignored. This intentionally keeps bonus answers short and forgiving.
    """
    return "".join(ch for ch in str(value or "").casefold() if ch.isalnum())


def short_answer_matches(submitted, accepted):
    normalized = normalize_short_answer(submitted)
    if not normalized:
        return False
    return any(normalized == normalize_short_answer(option) for option in str(accepted or "").split("|"))


def valid_student_section(program, class_section, subject_code="CSDC101"):
    """Validate the recognized class identifiers for each supported course."""
    program = str(program or "").strip().upper()
    class_section = str(class_section or "").strip().upper()
    subject_code = normalize_subject_code(subject_code)
    if subject_code == "CSEC303":
        return program in CSEC303_SECTIONS and class_section in CSEC303_SECTIONS[program]
    return program in STUDENT_SECTIONS and class_section in STUDENT_SECTIONS[program]


def parse_student_identity(form, subject_code="CSDC101"):
    """Validate the student identity fields shared by Home and Student View."""
    first_name = re.sub(r"\s+", " ", form.get("first_name", "").strip())
    last_name = re.sub(r"\s+", " ", form.get("last_name", "").strip())
    program = form.get("program", "").strip().upper()
    class_section = form.get("class_section", "").strip()
    if len(first_name) < 1 or len(first_name) > 60:
        return None, "Enter your first name."
    if len(last_name) < 1 or len(last_name) > 60:
        return None, "Enter your last name."
    name = f"{first_name} {last_name}".strip()
    if not valid_student_section(program, class_section, subject_code):
        if normalize_subject_code(subject_code) == "CSDC101":
            return None, "Choose a valid CSDC101 program and section."
        return None, "CSEC303 is assigned to section ZC32. Use Program/Class ZC and Section 32."
    return (first_name, last_name, name, program, class_section), None


def parse_iso(value):
    if not value:
        return None
    try:
        dt = datetime.fromisoformat(value)
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=APP_TZ)
        else:
            dt = dt.astimezone(APP_TZ)
        return dt
    except ValueError:
        return None


def csrf_token():
    token = session.get("csrf_token")
    if not token:
        token = secrets.token_urlsafe(32)
        session["csrf_token"] = token
    return token


app.jinja_env.globals["csrf_token"] = csrf_token


def _request_subject_code():
    """Resolve the course identity for assessment-specific pages.

    General pages stay neutral. Assessment pages derive their subject from the
    assessment/batch/session identifier in the current request so Instructor
    View and Student View use the same course header, background, and watermark.
    """
    endpoint = request.endpoint or ""
    legacy = str(request.args.get("assessment") or "").strip().lower()
    if endpoint == "admin_dashboard" and legacy in {"midterm", "posttest"}:
        return "CSDC101"
    view_args = request.view_args or {}
    assessment_id = view_args.get("assessment_id") or request.args.get("assessment_id")
    batch_id = view_args.get("batch_id")
    sid = view_args.get("sid")
    qid = view_args.get("qid")
    conn = None
    try:
        if assessment_id and str(assessment_id).isdigit():
            conn = connect()
            row = conn.execute(
                """SELECT s.code FROM assessments a JOIN subjects s ON s.id=a.subject_id WHERE a.id=?""",
                (int(assessment_id),),
            ).fetchone()
            return normalize_subject_code(row["code"]) if row and row["code"] else None
        if batch_id and str(batch_id).isdigit():
            conn = connect()
            row = conn.execute(
                """SELECT COALESCE(s.code,sa.code,'CSDC101') AS code
                   FROM batches b
                   LEFT JOIN subjects s ON s.id=b.subject_id
                   LEFT JOIN assessments a ON a.id=b.assessment_id
                   LEFT JOIN subjects sa ON sa.id=a.subject_id
                   WHERE b.id=?""",
                (int(batch_id),),
            ).fetchone()
            return normalize_subject_code(row["code"]) if row and row["code"] else None
        if sid and str(sid).isdigit() and endpoint.startswith("admin_"):
            conn = connect()
            row = conn.execute(
                """SELECT s.code FROM exam_sessions e JOIN assessments a ON a.id=e.assessment_id JOIN subjects s ON s.id=a.subject_id WHERE e.id=?""",
                (int(sid),),
            ).fetchone()
            return normalize_subject_code(row["code"]) if row and row["code"] else None
        if qid and str(qid).isdigit():
            conn = connect()
            row = conn.execute(
                """SELECT s.code FROM questions q JOIN assessments a ON a.id=q.assessment_id JOIN subjects s ON s.id=a.subject_id WHERE q.id=?""",
                (int(qid),),
            ).fetchone()
            return normalize_subject_code(row["code"]) if row and row["code"] else None
    except Exception:
        return None
    finally:
        if conn is not None:
            conn.close()
    return None


@app.context_processor
def inject_app_identity():
    return {
        "app_name": APP_NAME,
        "app_version": APP_VERSION,
        "app_release_species": APP_RELEASE_SPECIES,
        "app_release_scientific_name": APP_RELEASE_SCIENTIFIC_NAME,
        "app_release_common_name": APP_RELEASE_COMMON_NAME,
        "app_asset_revision": APP_ASSET_REVISION,
        "custos_theme": get_saved_theme(),
        "admin_assessment": session.get("admin_assessment", "posttest"),
        "selected_subject_code": normalize_subject_code(session.get("selected_subject_code", "CSDC101")),
        "supported_subjects": SUPPORTED_SUBJECTS,
        "page_subject_code": _request_subject_code(),
    }


def require_csrf():
    supplied = request.headers.get("X-CSRFToken") or request.form.get("csrf_token")
    if not supplied or not secrets.compare_digest(supplied, session.get("csrf_token", "")):
        abort(400, "Invalid CSRF token")


def admin_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        # Admin access requires an explicit instructor authentication event.
        # This prevents a stale/shared browser session from exposing Workspace
        # merely because an old admin_id remains in the signed cookie.
        if not session.get("admin_id") or session.get("admin_authenticated") is not True:
            return redirect(url_for("admin_login", next=request.full_path if request.query_string else request.path))
        conn = connect()
        try:
            admin = conn.execute("SELECT id,active FROM admins WHERE id=?", (session.get("admin_id"),)).fetchone()
        finally:
            conn.close()
        if not admin or ("active" in admin.keys() and not admin["active"]):
            for key in ("admin_id", "admin_role", "admin_display_name", "admin_auth_method", "admin_authenticated", "admin_authenticated_at", "owner_view_all"):
                session.pop(key, None)
            return redirect(url_for("admin_login"))
        return view(*args, **kwargs)
    return wrapped


def student_session_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        sid = session.get("student_session_id")
        if not sid:
            return redirect(url_for("student_login"))
        return view(*args, **kwargs)
    return wrapped


def batch_is_open(batch):
    now = datetime.now(APP_TZ)
    if not batch["active"]:
        return False, "This batch is currently disabled."
    open_at = parse_iso(batch["open_at"])
    close_at = parse_iso(batch["close_at"])
    if open_at and now < open_at:
        return False, f"The batch opens at {open_at.strftime('%Y-%m-%d %I:%M %p')}."
    if close_at and now > close_at:
        return False, "The allowed start window for this batch has closed."
    return True, None


def session_remaining(exam_session, batch):
    # Instructor trial sessions can be explicitly untimed. They remain open
    # until the instructor submits or deletes them.
    if "untimed" in exam_session.keys() and exam_session["untimed"]:
        return None, None
    started = parse_iso(exam_session["started_at"])
    deadline = started + timedelta(minutes=batch["duration_minutes"])
    # A delivery close time is also a hard deadline. This is important for
    # multi-day take-home assessments: a late start never extends the class due date.
    close_at = parse_iso(batch["close_at"]) if "close_at" in batch.keys() else None
    if close_at and close_at < deadline:
        deadline = close_at
    remaining = int((deadline - datetime.now(APP_TZ)).total_seconds())
    return max(0, remaining), deadline


def get_security_state(exam_session):
    """Return the server-authoritative security lock state for one attempt."""
    now = datetime.now(APP_TZ)
    temp_until = parse_iso(exam_session["temp_locked_until"]) if "temp_locked_until" in exam_session.keys() else None
    temp_remaining = 0
    if temp_until and temp_until > now:
        temp_remaining = max(0, math.ceil((temp_until - now).total_seconds()))
    return {
        "violation_count": int(exam_session["violation_count"] or 0) if "violation_count" in exam_session.keys() else 0,
        "permanent": bool(exam_session["security_locked"]) if "security_locked" in exam_session.keys() else False,
        "pending": bool(exam_session["pending_blackout"]) if "pending_blackout" in exam_session.keys() else False,
        "resume_required": bool(exam_session["security_resume_required"]) if "security_resume_required" in exam_session.keys() else False,
        "temp_remaining": temp_remaining,
    }


def session_security_mode(conn, exam_session):
    """Resolve the configured security mode for one student attempt."""
    if not conn or not exam_session:
        return "strict"
    try:
        row = conn.execute(
            """SELECT COALESCE(a.security_mode,'strict') AS security_mode
               FROM exam_sessions e
               LEFT JOIN batches b ON b.id=e.batch_id
               LEFT JOIN assessments a ON a.id=COALESCE(e.assessment_id,b.assessment_id)
               WHERE e.id=?""",
            (exam_session["id"],),
        ).fetchone()
        return (row["security_mode"] if row and row["security_mode"] else "strict").lower()
    except Exception:
        return "strict"


def session_subject_code(conn, exam_session):
    """Resolve the owning subject code for an attempt, defaulting legacy sessions to CSDC101."""
    if not conn or not exam_session:
        return "CSDC101"
    try:
        row = conn.execute(
            """SELECT COALESCE(s.code,'CSDC101') AS subject_code
               FROM exam_sessions e
               LEFT JOIN batches b ON b.id=e.batch_id
               LEFT JOIN assessments a ON a.id=COALESCE(e.assessment_id,b.assessment_id)
               LEFT JOIN subjects s ON s.id=COALESCE(a.subject_id,b.subject_id)
               WHERE e.id=?""",
            (exam_session["id"],),
        ).fetchone()
        return str(row["subject_code"] if row and row["subject_code"] else "CSDC101").upper()
    except Exception:
        return "CSDC101"


def clear_non_strict_security(conn, sid):
    """Clear lock-only state for Standard/Practice assessments."""
    conn.execute(
        """UPDATE exam_sessions
           SET security_locked=0, temp_locked_until=NULL, pending_blackout=0, security_resume_required=0
           WHERE id=?""",
        (sid,),
    )


def session_security_blocked(exam_session, conn=None):
    if conn is not None and session_security_mode(conn, exam_session) != "strict":
        state = get_security_state(exam_session)
        state.update({"permanent": False, "pending": False, "resume_required": False, "temp_remaining": 0})
        return False, state
    state = get_security_state(exam_session)
    return state["permanent"] or state["pending"] or state["resume_required"] or state["temp_remaining"] > 0, state


def assign_questions_to_session(conn, sid, batch_slot):
    """Assign questions for either legacy fixed-format exams or free-form custom assessments."""
    batch = conn.execute(
        "SELECT assessment_type,subject_id,assessment_id FROM batches WHERE slot=?", (batch_slot,)
    ).fetchone()
    if not batch:
        raise ValueError("Assessment delivery set was not found.")

    rng = random.SystemRandom()
    if batch["assessment_type"] == "custom" and batch["assessment_id"]:
        assessment = conn.execute("SELECT * FROM assessments WHERE id=?", (batch["assessment_id"],)).fetchone()
        qrows = conn.execute(
            """SELECT id FROM questions WHERE assessment_id=? AND COALESCE(active,1)=1
               ORDER BY COALESCE(position,id),id""",
            (batch["assessment_id"],),
        ).fetchall()
        selected = [r["id"] for r in qrows]
        if not selected:
            raise ValueError("This custom assessment does not have any active questions yet.")
        limit = int(assessment["question_limit"] or 0) if assessment and "question_limit" in assessment.keys() else 0
        if limit > 0:
            if len(selected) < limit:
                raise ValueError(f"This assessment is configured for {limit} questions but only {len(selected)} are active.")
            selected = rng.sample(selected, limit) if assessment["shuffle_questions"] else selected[:limit]
        elif assessment and assessment["shuffle_questions"]:
            rng.shuffle(selected)
        for q_order, qid in enumerate(selected, start=1):
            option_order = ["A", "B", "C", "D"]
            if not assessment or assessment["shuffle_options"]:
                rng.shuffle(option_order)
            conn.execute(
                "INSERT INTO session_questions(session_id,question_id,q_order,option_order) VALUES (?,?,?,?)",
                (sid, qid, q_order, json.dumps(option_order)),
            )
        return

    # Legacy CSDC101 Midterm/Post-test flow remains unchanged.
    qrows = conn.execute(
        "SELECT id, part FROM questions WHERE batch_slot=? AND COALESCE(active,1)=1 ORDER BY id",
        (batch_slot,),
    ).fetchall()
    p1 = [r["id"] for r in qrows if r["part"] == 1]
    p2 = [r["id"] for r in qrows if r["part"] == 2]
    if len(p1) < 40 or len(p2) < 20:
        raise ValueError(
            f"Question bank {batch_slot} needs at least 40 active Part I and 20 active Part II questions "
            f"(currently {len(p1)} and {len(p2)})."
        )
    selected_p1 = rng.sample(p1, 40)
    selected_p2 = rng.sample(p2, 20)
    rng.shuffle(selected_p1)
    rng.shuffle(selected_p2)
    for q_order, qid in enumerate(selected_p1 + selected_p2, start=1):
        option_order = ["A", "B", "C", "D"]
        rng.shuffle(option_order)
        conn.execute(
            "INSERT INTO session_questions(session_id,question_id,q_order,option_order) VALUES (?,?,?,?)",
            (sid, qid, q_order, json.dumps(option_order)),
        )

    if batch["assessment_type"] == "midterm":
        bonus = conn.execute(
            "SELECT id,position FROM bonus_questions WHERE assessment_type='midterm' AND active=1 ORDER BY position"
        ).fetchall()
        if len(bonus) != 5:
            raise ValueError(f"Midterm requires exactly 5 active bonus questions; found {len(bonus)}.")
        for b in bonus:
            conn.execute(
                """INSERT INTO session_bonus_answers(session_id,bonus_question_id,q_order,answer_text)
                   VALUES (?,?,?, '') ON CONFLICT(session_id,bonus_question_id) DO NOTHING""",
                (sid, b["id"], b["position"]),
            )


def finalize_exam(conn, session_id, reason="submitted"):
    exam = conn.execute("SELECT * FROM exam_sessions WHERE id=?", (session_id,)).fetchone()
    if not exam or exam["status"] == "submitted":
        return exam

    assessment = None
    if "assessment_id" in exam.keys() and exam["assessment_id"]:
        assessment = conn.execute("SELECT * FROM assessments WHERE id=?", (exam["assessment_id"],)).fetchone()

    rows = conn.execute(
        """SELECT q.part, q.correct_option, q.points, sq.selected_option
           FROM session_questions sq
           JOIN questions q ON q.id = sq.question_id
           WHERE sq.session_id=?""",
        (session_id,),
    ).fetchall()

    if assessment and assessment["assessment_type"] == "custom":
        correct = sum(1 for r in rows if r["selected_option"] == r["correct_option"])
        score = sum(int(r["points"] or 1) for r in rows if r["selected_option"] == r["correct_option"])
        p1_correct, p1_score = correct, score
        p2_correct = p2_score = 0
        bonus_correct = bonus_score = 0
        total = score
    else:
        p1_correct = sum(1 for r in rows if r["part"] == 1 and r["selected_option"] == r["correct_option"])
        p2_correct = sum(1 for r in rows if r["part"] == 2 and r["selected_option"] == r["correct_option"])
        p1_score = min(p1_correct, 30)
        p2_score = p2_correct
        bonus_rows = conn.execute(
            """SELECT sba.answer_text, b.accepted_answer
               FROM session_bonus_answers sba
               JOIN bonus_questions b ON b.id=sba.bonus_question_id
               WHERE sba.session_id=?""",
            (session_id,),
        ).fetchall()
        bonus_correct = sum(1 for r in bonus_rows if short_answer_matches(r["answer_text"], r["accepted_answer"]))
        bonus_score = bonus_correct
        total = p1_score + p2_score + bonus_score

    submitted_at = iso_now()
    conn.execute(
        """UPDATE exam_sessions
           SET submitted_at=?, status='submitted', part1_correct=?, part1_score=?,
               part2_correct=?, part2_score=?, bonus_correct=?, bonus_score=?, auto_total=?
           WHERE id=?""",
        (submitted_at, p1_correct, p1_score, p2_correct, p2_score, bonus_correct, bonus_score, total, session_id),
    )
    conn.execute(
        "INSERT INTO proctor_events(session_id,event_type,detail,created_at) VALUES (?,?,?,?)",
        (session_id, "exam_finalized", reason, submitted_at),
    )
    conn.commit()
    return conn.execute("SELECT * FROM exam_sessions WHERE id=?", (session_id,)).fetchone()


@app.post("/api/preferences/theme")
def save_theme_preference():
    require_csrf()
    payload = request.get_json(silent=True) or {}
    theme = str(payload.get("theme", "")).strip().lower()
    if theme not in {"light", "dark"}:
        return jsonify({"ok": False, "error": "Invalid theme."}), 400

    session["ui_theme"] = theme
    return jsonify({"ok": True, "theme": theme})


@app.after_request
def add_security_headers(resp):
    resp.headers["X-Frame-Options"] = "DENY"
    resp.headers["X-Content-Type-Options"] = "nosniff"
    resp.headers["Referrer-Policy"] = "no-referrer"
    resp.headers["Permissions-Policy"] = "camera=(), microphone=(), geolocation=(), payment=(), usb=(), serial=(), bluetooth=()"
    resp.headers["Cross-Origin-Opener-Policy"] = "same-origin-allow-popups"
    resp.headers["X-Permitted-Cross-Domain-Policies"] = "none"
    resp.headers["Origin-Agent-Cluster"] = "?1"
    if request.is_secure or request.headers.get("X-Forwarded-Proto", "").lower() == "https":
        resp.headers["Strict-Transport-Security"] = "max-age=31536000; includeSubDomains"
    g_csp = csp_additions()
    pyodide_csp = request.path == "/exam"
    pyodide_script = "https://cdn.jsdelivr.net https://cdnjs.cloudflare.com 'wasm-unsafe-eval'" if pyodide_csp else ""
    pyodide_connect = "https://cdn.jsdelivr.net https://cdnjs.cloudflare.com" if pyodide_csp else ""
    pyodide_worker = "worker-src 'self' blob:; " if pyodide_csp else ""
    resp.headers["Content-Security-Policy"] = (
        "default-src 'self'; img-src 'self' data: blob:; "
        f"style-src 'self' https://fonts.googleapis.com https://cdnjs.cloudflare.com {g_csp.get('style-src', '')}; "
        "font-src 'self' https://fonts.gstatic.com; "
        f"script-src 'self' {pyodide_script} {g_csp.get('script-src', '')}; "
        f"connect-src 'self' {pyodide_connect} {g_csp.get('connect-src', '')}; "
        + pyodide_worker
        + (f"frame-src {g_csp['frame-src']}; " if g_csp else "")
        + "frame-ancestors 'none'; base-uri 'self'; form-action 'self'"
    )
    if g_csp and request.endpoint in GOOGLE_SIGNIN_ENDPOINTS:
        # Google's sign-in button needs the page origin; send it only on these pages.
        resp.headers["Referrer-Policy"] = "strict-origin-when-cross-origin"

    # Versioned static assets are immutable for the lifetime of a release. The
    # old build marked every response no-store, forcing CSS/images/JS to download
    # again on repeat visits. Keep dynamic assessment/API responses uncached.
    if request.path.startswith("/static/"):
        if request.args.get("v"):
            resp.headers["Cache-Control"] = "public, max-age=31536000, immutable"
        else:
            # Unversioned URLs must never be held by a CDN: pages always use ?v=,
            # and a stale copy here was pre-cached by the service worker.
            resp.headers["Cache-Control"] = "no-cache"
        resp.headers.pop("Pragma", None)
    elif request.path == "/sw.js":
        resp.headers["Cache-Control"] = "no-cache, max-age=0"
        resp.headers["Pragma"] = "no-cache"
    elif request.path == "/manifest.webmanifest":
        resp.headers["Cache-Control"] = "public, max-age=3600"
        resp.headers.pop("Pragma", None)
    else:
        resp.headers["Cache-Control"] = "no-store, max-age=0"
        resp.headers["Pragma"] = "no-cache"
    return resp


@app.get("/manifest.webmanifest")
def pwa_manifest():
    return send_from_directory(
        app.static_folder,
        "manifest.webmanifest",
        mimetype="application/manifest+json",
    )


@app.get("/sw.js")
def pwa_service_worker():
    resp = send_from_directory(
        app.static_folder,
        "sw.js",
        mimetype="application/javascript",
    )
    resp.headers["Service-Worker-Allowed"] = "/"
    return resp


@app.get("/health")
def health():
    conn = connect()
    try:
        conn.execute("SELECT 1").fetchone()
        return jsonify({"status": "ok", "database": DATABASE_ENGINE}), 200
    finally:
        conn.close()


@app.route("/")
def index():
    return render_template("portal.html")



def _student_identity_from_session():
    """Return the durable signed-in student identity available to this browser."""
    ident = google_student_identity()
    if ident:
        return {
            "email": str(ident.get("email") or "").strip().lower(),
            "first_name": str(ident.get("given_name") or "").strip(),
            "last_name": str(ident.get("family_name") or "").strip(),
            "name": str(ident.get("name") or ident.get("email") or "").strip(),
            "verified": True,
        }
    stored = session.get("student_identity")
    if isinstance(stored, dict) and stored.get("email"):
        return stored
    return None


def _format_duration(minutes):
    try:
        minutes = int(minutes or 0)
    except (TypeError, ValueError):
        return "—"
    if minutes >= 1440 and minutes % 1440 == 0:
        days = minutes // 1440
        return f"{days} day{'s' if days != 1 else ''}"
    if minutes >= 60 and minutes % 60 == 0:
        hours = minutes // 60
        return f"{hours} hour{'s' if hours != 1 else ''}"
    return f"{minutes} min"

def _clear_instructor_identity():
    """Remove instructor authority when this browser explicitly enters Student View."""
    for key in (
        "admin_id", "admin_role", "admin_display_name", "admin_auth_method",
        "admin_authenticated", "admin_authenticated_at", "owner_view_all",
    ):
        session.pop(key, None)




@app.get("/student")
def student_classroom():
    """Student assessment home that can rediscover attempts across devices/sign-ins."""
    ident = _student_identity_from_session()
    if not ident:
        session["student_after_signin"] = url_for("student_classroom")
        return redirect(url_for("student_login"))
    # Student and instructor authority must never coexist in the same normal
    # browser session. Instructor preview routes do not use /student.
    _clear_instructor_identity()
    email = str(ident.get("email") or "").lower()
    conn = connect()
    # Any assessment explicitly assigned through Google Classroom, plus every
    # assessment this student has already started. This makes the dashboard
    # useful even after a roster is later refreshed or removed.
    rows = conn.execute(
        """SELECT DISTINCT a.id,a.title,a.display_type,a.assessment_type,a.description,a.duration_minutes,
                  a.start_at,a.end_at,a.active,a.security_mode,a.reveal_score,a.reveal_answers,
                  s.code AS subject_code,s.name AS subject_name,
                  r.course_name,r.course_section,r.program AS roster_program,r.class_section AS roster_section
           FROM assessments a
           JOIN subjects s ON s.id=a.subject_id
           LEFT JOIN classroom_rosters r ON r.assessment_id=a.id
           LEFT JOIN classroom_roster_students rs ON rs.roster_id=r.id AND LOWER(rs.email)=LOWER(?)
           WHERE a.deleted_at IS NULL
             AND (rs.id IS NOT NULL OR EXISTS (SELECT 1 FROM exam_sessions e WHERE e.assessment_id=a.id AND LOWER(e.email)=LOWER(?))
                  OR EXISTS (SELECT 1 FROM coding_sessions cs JOIN programming_labs pl ON pl.id=cs.lab_id WHERE pl.assessment_id=a.id AND LOWER(cs.email)=LOWER(?)))
           ORDER BY CASE WHEN a.end_at IS NULL THEN 1 ELSE 0 END,a.end_at DESC,a.id DESC""",
        (email,email,email),
    ).fetchall()
    cards=[]
    now=datetime.now(APP_TZ)
    for row in rows:
        a=dict(row)
        ex=conn.execute(
            "SELECT * FROM exam_sessions WHERE assessment_id=? AND LOWER(email)=LOWER(?) ORDER BY id DESC LIMIT 1",
            (a["id"],email),
        ).fetchone()
        coding=None
        if a["assessment_type"]=="programming_lab":
            coding=conn.execute(
                """SELECT cs.* FROM coding_sessions cs JOIN programming_labs pl ON pl.id=cs.lab_id
                   WHERE pl.assessment_id=? AND LOWER(cs.email)=LOWER(?) ORDER BY cs.id DESC LIMIT 1""",
                (a["id"],email),
            ).fetchone()
        batches=conn.execute("SELECT * FROM batches WHERE assessment_id=? AND active=1 ORDER BY id",(a["id"],)).fetchall()
        status="available"
        status_label="Available"
        action="open"
        if ex:
            if ex["status"]=="submitted": status,status_label,action="submitted","Submitted","result"
            else: status,status_label,action="in_progress","In progress","resume"
        elif coding:
            if coding["status"]=="submitted": status,status_label,action="submitted","Submitted","coding_result"
            else: status,status_label,action="in_progress","In progress","coding_resume"
        start=parse_iso(a.get("start_at")); end=parse_iso(a.get("end_at"))
        if not ex and not coding:
            if not a["active"]: status,status_label="closed","Unavailable"
            elif start and now < start: status,status_label="upcoming","Upcoming"
            elif end and now > end: status,status_label="closed","Closed"
        a.update({
            "attempt": dict(ex) if ex else None,
            "coding_attempt": dict(coding) if coding else None,
            "batch_count": len(batches),
            "status":status,"status_label":status_label,"action":action,
            "duration_label":_format_duration(a.get("duration_minutes")),
        })
        cards.append(a)
    conn.close()
    groups={"in_progress":[],"available":[],"upcoming":[],"submitted":[],"closed":[]}
    for c in cards: groups.setdefault(c["status"],[]).append(c)
    return render_template("student_classroom.html", identity=ident, assessment_groups=groups)


@app.get("/student/assessment/<int:assessment_id>/open")
def student_assessment_open(assessment_id):
    ident=_student_identity_from_session()
    if not ident:
        session["student_after_signin"] = url_for("student_assessment_open", assessment_id=assessment_id)
        return redirect(url_for("student_login"))
    email=str(ident.get("email") or "").lower()
    conn=connect()
    assessment=conn.execute(
        """SELECT a.*,s.code AS subject_code FROM assessments a JOIN subjects s ON s.id=a.subject_id
           WHERE a.id=? AND a.deleted_at IS NULL""",(assessment_id,)
    ).fetchone()
    if not assessment:
        conn.close(); abort(404)
    # Existing attempts always resume by server-side identity, independent of browser cookie.
    existing=conn.execute("SELECT * FROM exam_sessions WHERE assessment_id=? AND LOWER(email)=LOWER(?) ORDER BY id DESC LIMIT 1",(assessment_id,email)).fetchone()
    if existing:
        session.permanent=True
        session["student_session_id"]=existing["id"]
        conn.close()
        return redirect(url_for("result" if existing["status"]=="submitted" else "exam"))
    # Programming Labs continue through their established IDE login workflow.
    if assessment["assessment_type"]=="programming_lab":
        conn.close(); return redirect(url_for("nextgen.ide_login"))
    has_roster, rostered=roster_entry(conn,assessment_id,email)
    if not has_roster or not rostered:
        conn.close(); flash("This assessment is not assigned to your signed-in Classroom account.","error")
        return redirect(url_for("student_classroom"))
    batches=conn.execute("SELECT * FROM batches WHERE assessment_id=? AND active=1 ORDER BY id",(assessment_id,)).fetchall()
    if len(batches)!=1:
        conn.close(); flash("This assessment uses multiple delivery sets. Enter the session key issued for your set.","info")
        return redirect(url_for("student_login",subject=assessment["subject_code"]))
    batch=batches[0]
    is_open,message=batch_is_open(batch)
    if not is_open:
        conn.close(); flash(message or "This assessment is not currently open.","error")
        return redirect(url_for("student_classroom"))
    first_name=rostered["first_name"] or ident.get("first_name") or "Student"
    last_name=rostered["last_name"] or ident.get("last_name") or "-"
    student_name=f"{first_name} {last_name}".strip()
    session.permanent=True
    session["pending_email"]=email
    session["pending_first_name"]=first_name
    session["pending_last_name"]=last_name
    session["pending_student_name"]=student_name
    if str(assessment["subject_code"] or "").upper() == "CSEC303":
        session["pending_program"] = "ZC"
        session["pending_class_section"] = "32"
    else:
        session["pending_program"] = rostered["program"]
        session["pending_class_section"] = rostered["class_section"]
    session["pending_batch_id"]=batch["id"]
    session["pending_session_key"]=batch["access_code"]
    gident=google_student_identity()
    session["pending_google_sub"]=(gident or {}).get("sub","")
    session["student_identity"]={"email":email,"first_name":first_name,"last_name":last_name,"name":student_name,"verified":bool(gident)}
    conn.close()
    return redirect(url_for("instructions"))


@app.route("/login", methods=["GET", "POST"])
def student_login():
    """Identity-first student entry.

    Students authenticate first, then land on My Assessments. Session keys are
    intentionally entered from that dashboard rather than being mixed into the
    sign-in step. The same endpoint still accepts a session_key POST from the
    dashboard so existing Google Classroom and resume logic stays centralized.
    """
    if request.method == "GET":
        if (google_student_identity() or session.get("student_identity")) and request.args.get("force") != "1":
            return redirect(url_for("student_classroom"))
        return render_template("student_login.html", domain=ALLOWED_EMAIL_DOMAIN)

    require_csrf()
    google_ident = google_student_identity()
    session_key = request.form.get("session_key", "").strip().upper()

    # Sign-in only. Google normally reaches this state through /auth/google/student,
    # while the manual fallback is retained for local/offline testing.
    if not session_key:
        if STUDENT_GOOGLE_LOGIN_REQUIRED and not google_ident:
            flash("Sign in with your school Google account first.", "error")
            return render_template("student_login.html", domain=ALLOWED_EMAIL_DOMAIN)
        if google_ident:
            _clear_instructor_identity()
            first_name, last_name = google_profile_names(google_ident)
            email = google_ident["email"]
            session.permanent = True
            session["student_identity"] = {
                "email": email, "first_name": first_name, "last_name": last_name,
                "name": f"{first_name} {last_name}".strip(), "verified": True,
            }
            return redirect(url_for("student_classroom"))
        first_name = re.sub(r"\s+", " ", request.form.get("first_name", "").strip())[:60]
        last_name = re.sub(r"\s+", " ", request.form.get("last_name", "").strip())[:60]
        email = request.form.get("email", "").strip().lower()
        email_re = rf"^[A-Za-z0-9._%+\-]+@{re.escape(ALLOWED_EMAIL_DOMAIN)}$"
        if not first_name or not last_name or not re.match(email_re, email):
            flash(f"Enter your name and use your @{ALLOWED_EMAIL_DOMAIN} account.", "error")
            return render_template("student_login.html", domain=ALLOWED_EMAIL_DOMAIN)
        _clear_instructor_identity()
        session.permanent = True
        session["student_identity"] = {
            "email": email, "first_name": first_name, "last_name": last_name,
            "name": f"{first_name} {last_name}".strip(), "verified": False,
        }
        return redirect(url_for("student_classroom"))

    selected_subject = normalize_subject_code(request.form.get("subject_code") or session.get("selected_subject_code"))
    session["selected_subject_code"] = selected_subject
    ident = _student_identity_from_session()
    if STUDENT_GOOGLE_LOGIN_REQUIRED and not google_ident:
        flash("Sign in with your school Google account first.", "error")
        return redirect(url_for("student_login", force=1))
    if not ident:
        flash("Sign in before entering a session key.", "error")
        return redirect(url_for("student_login", force=1))

    email = str((google_ident or {}).get("email") or ident.get("email") or "").lower()
    if google_ident:
        first_name, last_name = google_profile_names(google_ident)
    else:
        first_name = str(ident.get("first_name") or "Student")[:60]
        last_name = str(ident.get("last_name") or "-")[:60]
    program = request.form.get("program", "").strip().upper()
    class_section = request.form.get("class_section", "").strip()

    conn = connect()
    batch = conn.execute("SELECT * FROM batches WHERE access_code=?", (session_key,)).fetchone()
    if not batch:
        conn.close()
        flash("Invalid session key. Check the key announced by your instructor.", "error")
        return redirect(url_for("student_classroom"))

    assessment = None
    actual_subject_code = "CSDC101"
    if "assessment_id" in batch.keys() and batch["assessment_id"]:
        assessment = conn.execute(
            """SELECT a.*, s.code AS subject_code, s.name AS subject_name
               FROM assessments a JOIN subjects s ON s.id=a.subject_id WHERE a.id=?""",
            (batch["assessment_id"],),
        ).fetchone()
        if assessment and assessment["subject_code"]:
            actual_subject_code = str(assessment["subject_code"]).upper()
    elif "subject_id" in batch.keys() and batch["subject_id"]:
        srow = conn.execute("SELECT code FROM subjects WHERE id=?", (batch["subject_id"],)).fetchone()
        if srow and srow["code"]:
            actual_subject_code = str(srow["code"]).upper()
    if selected_subject and actual_subject_code != selected_subject:
        conn.close()
        flash(f"That session key belongs to {actual_subject_code}. Change the subject and try again.", "error")
        return redirect(url_for("student_classroom"))

    session["selected_subject_code"] = actual_subject_code
    session.permanent = True
    has_roster, rostered = roster_entry(conn, assessment["id"] if assessment else None, email)
    if has_roster:
        if not rostered:
            conn.close()
            flash("Your Google account is not on the class list for this assessment. Ask your instructor.", "error")
            return redirect(url_for("student_classroom"))
        first_name = rostered["first_name"] or first_name
        last_name = rostered["last_name"] or last_name
        program, class_section = rostered["program"], rostered["class_section"]
    elif not valid_student_section(program, class_section, actual_subject_code):
        conn.close()
        flash("This assessment is not linked to a Classroom roster. Add your class details under Session Key and try again.", "error")
        return redirect(url_for("student_classroom"))

    if actual_subject_code == "CSEC303":
        program, class_section = "ZC", "32"

    student_name = f"{first_name} {last_name}".strip()
    session["student_identity"] = {
        "email": email, "first_name": first_name, "last_name": last_name,
        "name": student_name, "verified": bool(google_ident or ident.get("verified")),
    }
    if assessment and assessment["assessment_type"] == "custom":
        allowed = {x.strip().upper() for x in str(assessment["allowed_sections"] or "").split(",") if x.strip()}
        student_section = f"{program}{class_section}"
        if allowed and student_section not in allowed:
            conn.close()
            flash(f"This assessment is not assigned to {student_section}.", "error")
            return redirect(url_for("student_classroom"))

    existing = conn.execute("SELECT * FROM exam_sessions WHERE email=? AND batch_id=?", (email, batch["id"])).fetchone()
    if existing:
        conn.execute(
            "UPDATE exam_sessions SET first_name=?, last_name=?, student_name=?, program=?, class_section=?, monitor_done=0 WHERE id=?",
            (first_name, last_name, student_name, program, class_section, existing["id"]),
        )
        conn.commit()
        session["student_session_id"] = existing["id"]
        if existing["status"] == "submitted":
            conn.close(); return redirect(url_for("result"))
        if existing["terms_accepted_at"]:
            conn.close(); return redirect(url_for("exam"))

    is_open, message = batch_is_open(batch)
    if not is_open:
        conn.close(); flash(message, "error"); return redirect(url_for("student_classroom"))

    session["pending_email"] = email
    session["pending_first_name"] = first_name
    session["pending_last_name"] = last_name
    session["pending_student_name"] = student_name
    session["pending_program"] = program
    session["pending_class_section"] = class_section
    session["pending_batch_id"] = batch["id"]
    session["pending_session_key"] = session_key
    session["pending_google_sub"] = google_ident["sub"] if google_ident else ""
    conn.close()
    return redirect(url_for("instructions"))


@app.route("/instructions")
def instructions():
    email = session.get("pending_email")
    batch_id = session.get("pending_batch_id")
    if not email or not batch_id:
        return redirect(url_for("student_login"))
    conn = connect()
    batch = conn.execute(
        """SELECT b.*, COALESCE(s.code, sa.code, 'CSDC101') AS subject_code
           FROM batches b
           LEFT JOIN subjects s ON s.id=b.subject_id
           LEFT JOIN assessments a ON a.id=b.assessment_id
           LEFT JOIN subjects sa ON sa.id=a.subject_id
           WHERE b.id=?""",
        (batch_id,),
    ).fetchone()
    assessment = None
    custom_question_count = 0
    custom_max_score = 0
    if batch and "assessment_id" in batch.keys() and batch["assessment_id"]:
        assessment = conn.execute(
            """SELECT a.*,s.code AS subject_code,s.name AS subject_name
               FROM assessments a JOIN subjects s ON s.id=a.subject_id WHERE a.id=?""",
            (batch["assessment_id"],),
        ).fetchone()
        if assessment and assessment["assessment_type"] == "custom":
            qrows = conn.execute(
                "SELECT points FROM questions WHERE assessment_id=? AND COALESCE(active,1)=1 ORDER BY COALESCE(position,id),id",
                (assessment["id"],),
            ).fetchall()
            limit = int(assessment["question_limit"] or 0)
            if limit > 0:
                qrows = qrows[:limit]
            custom_question_count = len(qrows)
            custom_max_score = sum(int(r["points"] or 1) for r in qrows)
    conn.close()
    if not batch:
        return redirect(url_for("student_login"))
    return render_template(
        "instructions.html",
        batch=batch,
        assessment=assessment,
        custom_question_count=custom_question_count,
        custom_max_score=custom_max_score,
        email=email,
        first_name=session.get("pending_first_name"),
        last_name=session.get("pending_last_name"),
        student_name=session.get("pending_student_name"),
        program=session.get("pending_program"),
        class_section=session.get("pending_class_section"),
        duration_label=_format_duration(batch["duration_minutes"]),
    )


@app.route("/start", methods=["POST"])
def start_exam():
    require_csrf()
    email = session.get("pending_email")
    first_name = session.get("pending_first_name")
    last_name = session.get("pending_last_name")
    student_name = session.get("pending_student_name")
    program = session.get("pending_program")
    class_section = session.get("pending_class_section")
    batch_id = session.get("pending_batch_id")
    pending_session_key = session.get("pending_session_key", "")
    if not email or not first_name or not last_name or not student_name or not program or not class_section or not batch_id or not pending_session_key:
        return redirect(url_for("student_login"))
    if request.form.get("terms_accept") != "yes":
        flash("You must accept the assessment monitoring and integrity terms before starting.", "error")
        return redirect(url_for("instructions"))
    google_ident = google_student_identity()
    if STUDENT_GOOGLE_LOGIN_REQUIRED and (not google_ident or google_ident["email"] != email):
        flash("Sign in with your school Google account first.", "error")
        return redirect(url_for("student_login"))
    pending_google_sub = session.get("pending_google_sub", "")

    conn = connect()
    batch = conn.execute(
        """SELECT b.*, COALESCE(s.code, sa.code, 'CSDC101') AS subject_code
           FROM batches b
           LEFT JOIN subjects s ON s.id=b.subject_id
           LEFT JOIN assessments a ON a.id=b.assessment_id
           LEFT JOIN subjects sa ON sa.id=a.subject_id
           WHERE b.id=?""",
        (batch_id,),
    ).fetchone()
    if not batch:
        conn.close()
        abort(404)
    if not secrets.compare_digest(batch["access_code"], pending_session_key):
        conn.close()
        session.pop("pending_email", None)
        session.pop("pending_first_name", None)
        session.pop("pending_last_name", None)
        session.pop("pending_student_name", None)
        session.pop("pending_program", None)
        session.pop("pending_class_section", None)
        session.pop("pending_batch_id", None)
        session.pop("pending_session_key", None)
        flash("That session key has expired or was regenerated. Enter the current key to continue.", "error")
        return redirect(url_for("student_login"))
    is_open, message = batch_is_open(batch)
    if not is_open:
        conn.close()
        flash(message, "error")
        return redirect(url_for("student_login"))

    existing = conn.execute(
        "SELECT * FROM exam_sessions WHERE email=? AND batch_id=?", (email, batch_id)
    ).fetchone()
    accepted_at = iso_now()
    if existing:
        sid = existing["id"]
        conn.execute("UPDATE exam_sessions SET terms_accepted_at=COALESCE(terms_accepted_at, ?) WHERE id=?", (accepted_at, sid))
        conn.commit()
    else:
        cur = conn.execute(
            """INSERT INTO exam_sessions(email,first_name,last_name,student_name,program,class_section,batch_id,assessment_id,started_at,ip_address,user_agent,terms_accepted_at,auth_method,google_sub)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?) RETURNING id""",
            (email, first_name, last_name, student_name, program, class_section, batch_id, batch["assessment_id"], accepted_at, request.remote_addr, request.headers.get("User-Agent", "")[:500], accepted_at, "google" if pending_google_sub else "form", pending_google_sub or None),
        )
        sid = cur.fetchone()[0]
        try:
            assign_questions_to_session(conn, sid, batch["slot"])
        except ValueError as exc:
            conn.execute("DELETE FROM exam_sessions WHERE id=?", (sid,))
            conn.commit()
            conn.close()
            flash(f"This question bank is not ready: {exc}", "error")
            return redirect(url_for("student_login"))
        conn.execute(
            "INSERT INTO proctor_events(session_id,event_type,detail,created_at) VALUES (?,?,?,?)",
            (sid, "exam_started", "Student started the exam", iso_now()),
        )
        conn.commit()

    conn.close()
    session.pop("pending_email", None)
    session.pop("pending_first_name", None)
    session.pop("pending_last_name", None)
    session.pop("pending_student_name", None)
    session.pop("pending_program", None)
    session.pop("pending_class_section", None)
    session.pop("pending_batch_id", None)
    session.pop("pending_session_key", None)
    session.pop("pending_google_sub", None)
    session.permanent = True
    session["student_session_id"] = sid
    session["student_identity"] = {"email": email, "first_name": first_name, "last_name": last_name, "name": student_name, "verified": bool(pending_google_sub)}
    return redirect(url_for("exam"))


@app.route("/exam")
@student_session_required
def exam():
    sid = session["student_session_id"]
    conn = connect()
    exam_session = conn.execute("SELECT * FROM exam_sessions WHERE id=?", (sid,)).fetchone()
    if not exam_session:
        conn.close()
        session.pop("student_session_id", None)
        return redirect(url_for("student_login"))
    if exam_session["status"] == "submitted":
        conn.close()
        return redirect(url_for("result"))

    batch = conn.execute("SELECT * FROM batches WHERE id=?", (exam_session["batch_id"],)).fetchone()
    assessment = None
    if batch and "assessment_id" in batch.keys() and batch["assessment_id"]:
        assessment = conn.execute(
            """SELECT a.*,s.code AS subject_code,s.name AS subject_name
               FROM assessments a JOIN subjects s ON s.id=a.subject_id WHERE a.id=?""",
            (batch["assessment_id"],),
        ).fetchone()
    remaining, deadline = session_remaining(exam_session, batch)
    if remaining is not None and remaining <= 0:
        finalize_exam(conn, sid, "time_limit_reached")
        conn.close()
        return redirect(url_for("result"))

    lab_workbench = get_assessment_workbench(conn, assessment) if assessment else None
    workbench_aliases = {}
    workbench_sections_by_key = {}
    if lab_workbench:
        for section in lab_workbench.get("sections", []):
            key = str(section.get("key") or "").strip().upper()
            if key:
                workbench_aliases[re.sub(r"[^A-Z0-9]+", "", key)] = key
                workbench_sections_by_key[key] = section
            for alias in (section.get("title"), section.get("tab")):
                normalized = re.sub(r"[^A-Z0-9]+", "", str(alias or "").upper())
                if normalized:
                    workbench_aliases[normalized] = key

    rows = conn.execute(
        """SELECT sq.q_order, sq.option_order, sq.selected_option, COALESCE(sq.marked_for_review,0) AS marked_for_review,
                  q.id AS question_id, q.part, q.topic, q.prompt, q.code, q.points, q.workbench_section,
                  q.option_a, q.option_b, q.option_c, q.option_d
           FROM session_questions sq
           JOIN questions q ON q.id=sq.question_id
           WHERE sq.session_id=? ORDER BY sq.q_order""",
        (sid,),
    ).fetchall()
    questions = []
    for r in rows:
        base = {k: r[k] for k in r.keys()}
        option_map = {"A": r["option_a"], "B": r["option_b"], "C": r["option_c"], "D": r["option_d"]}
        base["display_options"] = [(letter, option_map[origin]) for letter, origin in zip("ABCD", json.loads(r["option_order"]))]
        # Preserve original option key in a separate tuple value for grading.
        base["display_options"] = [(display_letter, origin, option_map[origin]) for display_letter, origin in zip("ABCD", json.loads(r["option_order"]))]
        base["case_key"] = str(r["workbench_section"] or "").strip().upper() or "GUIDE"
        if not r["workbench_section"] and workbench_aliases:
            topic_normalized = re.sub(r"[^A-Z0-9]+", "", str(r["topic"] or "").upper())
            # Exact topic/section matches first, then safe prefix matches. This makes
            # ordinary notebooks work without hard-coding PHANTOM section names.
            if topic_normalized in workbench_aliases:
                base["case_key"] = workbench_aliases[topic_normalized]
            else:
                for alias, section_key in sorted(workbench_aliases.items(), key=lambda kv: len(kv[0]), reverse=True):
                    if alias and (topic_normalized.startswith(alias) or alias.startswith(topic_normalized)):
                        base["case_key"] = section_key
                        break
        linked_section = workbench_sections_by_key.get(base["case_key"])
        if linked_section and linked_section.get("preview"):
            base["media_url"] = linked_section["preview"]
            base["media_caption"] = linked_section.get("title") or "Workbench evidence preview"
        questions.append(base)
    bonus_questions = []
    if batch["assessment_type"] == "midterm":
        bonus_rows = conn.execute(
            """SELECT b.id AS bonus_question_id, b.position, b.topic, b.prompt, sba.answer_text
               FROM session_bonus_answers sba
               JOIN bonus_questions b ON b.id=sba.bonus_question_id
               WHERE sba.session_id=? ORDER BY sba.q_order""",
            (sid,),
        ).fetchall()
        bonus_questions = [dict(r) for r in bonus_rows]
    security = get_security_state(exam_session)
    total_count = len(questions) + len(bonus_questions)
    try:
        stored_position = int(exam_session["last_question_index"] or 0)
    except (KeyError, TypeError, ValueError):
        stored_position = 0
    resume_index = max(0, min(stored_position, max(total_count - 1, 0)))
    conn.close()
    return render_template(
        "exam.html", exam_session=exam_session, batch=batch, assessment=assessment, questions=questions,
        bonus_questions=bonus_questions, total_count=total_count, resume_index=resume_index,
        remaining=remaining, deadline=deadline.isoformat() if deadline else None,
        security=security, is_custom=bool(assessment and assessment["assessment_type"] == "custom"),
        security_mode=(assessment["security_mode"] if assessment else "strict"),
        subject_code=(assessment["subject_code"] if assessment and "subject_code" in assessment.keys() else "CSDC101"),
        lab_toolkit_sections=(lab_workbench["sections"] if lab_workbench else []),
        lab_workbench=lab_workbench
    )


@app.route("/api/question-position", methods=["POST"])
@student_session_required
def save_question_position():
    require_csrf()
    sid = session["student_session_id"]
    data = request.get_json(silent=True) or {}
    try:
        index = int(data.get("index", 0))
    except (TypeError, ValueError):
        return jsonify({"ok": False, "error": "invalid_index"}), 400
    if index < 0 or index > 200:
        return jsonify({"ok": False, "error": "invalid_index"}), 400

    conn = connect()
    ex = conn.execute("SELECT * FROM exam_sessions WHERE id=?", (sid,)).fetchone()
    if not ex:
        conn.close()
        return jsonify({"ok": False, "error": "session_not_found"}), 404
    if ex["status"] == "submitted":
        conn.close()
        return jsonify({"ok": False, "error": "submitted"}), 409
    blocked, security = session_security_blocked(ex, conn)
    if blocked:
        conn.close()
        return jsonify({"ok": False, "locked": True, **security}), 423
    conn.execute("UPDATE exam_sessions SET last_question_index=? WHERE id=?", (index, sid))
    conn.commit()
    conn.close()
    return jsonify({"ok": True, "index": index})


@app.route("/api/question-review", methods=["POST"])
@student_session_required
def set_question_review_flag():
    require_csrf()
    sid = session["student_session_id"]
    data = request.get_json(silent=True) or {}
    try:
        qid = int(data.get("question_id"))
    except (TypeError, ValueError):
        return jsonify({"ok": False, "error": "Invalid question"}), 400
    flagged = 1 if bool(data.get("flagged")) else 0

    conn = connect()
    ex = conn.execute("SELECT * FROM exam_sessions WHERE id=?", (sid,)).fetchone()
    if not ex:
        conn.close()
        return jsonify({"ok": False, "error": "session_not_found"}), 404
    if ex["status"] != "in_progress":
        conn.close()
        return jsonify({"ok": False, "error": "Exam is not active"}), 409
    blocked, security = session_security_blocked(ex, conn)
    if blocked:
        conn.close()
        return jsonify({"ok": False, "locked": True, **security}), 423
    cur = conn.execute(
        "UPDATE session_questions SET marked_for_review=? WHERE session_id=? AND question_id=?",
        (flagged, sid, qid),
    )
    conn.commit()
    conn.close()
    if cur.rowcount != 1:
        return jsonify({"ok": False, "error": "Question not assigned to session"}), 404
    return jsonify({"ok": True, "flagged": bool(flagged)})


@app.route("/api/answer", methods=["POST"])
@student_session_required
def save_answer():
    require_csrf()
    sid = session["student_session_id"]
    data = request.get_json(silent=True) or {}
    try:
        qid = int(data.get("question_id"))
    except (TypeError, ValueError):
        return jsonify({"ok": False, "error": "Invalid question"}), 400
    answer = str(data.get("answer", ""))
    if answer not in {"A", "B", "C", "D"}:
        return jsonify({"ok": False, "error": "Invalid answer"}), 400

    conn = connect()
    ex = conn.execute("SELECT * FROM exam_sessions WHERE id=?", (sid,)).fetchone()
    if not ex or ex["status"] != "in_progress":
        conn.close()
        return jsonify({"ok": False, "error": "Exam is not active"}), 409
    blocked, security = session_security_blocked(ex, conn)
    if blocked:
        conn.close()
        return jsonify({"ok": False, "locked": True, **security}), 423
    batch = conn.execute("SELECT * FROM batches WHERE id=?", (ex["batch_id"],)).fetchone()
    remaining, _ = session_remaining(ex, batch)
    if remaining is not None and remaining <= 0:
        finalize_exam(conn, sid, "time_limit_reached")
        conn.close()
        return jsonify({"ok": False, "expired": True}), 409

    cur = conn.execute(
        """UPDATE session_questions SET selected_option=?
           WHERE session_id=? AND question_id=?""",
        (answer, sid, qid),
    )
    conn.commit()
    conn.close()
    if cur.rowcount != 1:
        return jsonify({"ok": False, "error": "Question not assigned to session"}), 404
    return jsonify({"ok": True})


@app.route("/api/bonus-answer", methods=["POST"])
@student_session_required
def save_bonus_answer():
    require_csrf()
    sid = session["student_session_id"]
    data = request.get_json(silent=True) or {}
    try:
        qid = int(data.get("bonus_question_id"))
    except (TypeError, ValueError):
        return jsonify({"ok": False, "error": "Invalid bonus question"}), 400
    answer = str(data.get("answer", ""))[:120]

    conn = connect()
    ex = conn.execute("SELECT * FROM exam_sessions WHERE id=?", (sid,)).fetchone()
    if not ex or ex["status"] != "in_progress":
        conn.close()
        return jsonify({"ok": False, "error": "Exam is not active"}), 409
    blocked, security = session_security_blocked(ex, conn)
    if blocked:
        conn.close()
        return jsonify({"ok": False, "locked": True, **security}), 423
    batch = conn.execute("SELECT * FROM batches WHERE id=?", (ex["batch_id"],)).fetchone()
    remaining, _ = session_remaining(ex, batch)
    if remaining is not None and remaining <= 0:
        finalize_exam(conn, sid, "time_limit_reached")
        conn.close()
        return jsonify({"ok": False, "expired": True}), 409
    cur = conn.execute(
        "UPDATE session_bonus_answers SET answer_text=? WHERE session_id=? AND bonus_question_id=?",
        (answer, sid, qid),
    )
    conn.commit()
    conn.close()
    if cur.rowcount != 1:
        return jsonify({"ok": False, "error": "Bonus question not assigned to session"}), 404
    return jsonify({"ok": True})


@app.route("/api/proctor-events", methods=["POST"])
@student_session_required
def proctor_events():
    """Store low-volume proctor telemetry in batches.

    The browser buffers ordinary behavior events so a classroom does not open a
    database connection for every harmless click/shortcut. Security violations
    still use their dedicated immediate endpoint.
    """
    require_csrf()
    sid = session["student_session_id"]
    data = request.get_json(silent=True) or {}
    raw_events = data.get("events") or []
    if not isinstance(raw_events, list):
        return jsonify({"ok": False, "error": "events must be a list"}), 400
    raw_events = raw_events[:25]

    events = []
    flagged_delta = 0
    now = iso_now()
    for raw in raw_events:
        if not isinstance(raw, dict):
            continue
        event_type = str(raw.get("type", "unknown"))[:80]
        detail = str(raw.get("detail", ""))[:500]
        if not event_type:
            continue
        events.append((sid, event_type, detail, now))
        if event_type in SUSPICIOUS_EVENTS:
            flagged_delta += 1

    if not events:
        return jsonify({"ok": True, "stored": 0})

    conn = connect()
    ex = conn.execute("SELECT * FROM exam_sessions WHERE id=?", (sid,)).fetchone()
    if ex and ex["status"] == "in_progress":
        if session_security_mode(conn, ex) != "strict":
            clear_non_strict_security(conn, sid)
            conn.commit()
            conn.close()
            return jsonify({"ok": True, "ignored": True, "stored": 0})
        conn.executemany(
            "INSERT INTO proctor_events(session_id,event_type,detail,created_at) VALUES (?,?,?,?)",
            events,
        )
        if flagged_delta:
            conn.execute(
                "UPDATE exam_sessions SET flagged_count=flagged_count+? WHERE id=?",
                (flagged_delta, sid),
            )
        conn.commit()
        stored = len(events)
    else:
        stored = 0
    conn.close()
    return jsonify({"ok": True, "stored": stored})


@app.route("/api/proctor-event", methods=["POST"])
@student_session_required
def proctor_event():
    """Compatibility endpoint for older clients; new clients use batching."""
    require_csrf()
    sid = session["student_session_id"]
    data = request.get_json(silent=True) or {}
    event_type = str(data.get("type", "unknown"))[:80]
    detail = str(data.get("detail", ""))[:500]
    conn = connect()
    ex = conn.execute("SELECT * FROM exam_sessions WHERE id=?", (sid,)).fetchone()
    if ex and ex["status"] == "in_progress":
        if session_security_mode(conn, ex) != "strict":
            clear_non_strict_security(conn, sid)
            conn.commit()
            conn.close()
            return jsonify({"ok": True, "ignored": True})
        conn.execute(
            "INSERT INTO proctor_events(session_id,event_type,detail,created_at) VALUES (?,?,?,?)",
            (sid, event_type, detail, iso_now()),
        )
        if event_type in SUSPICIOUS_EVENTS:
            conn.execute("UPDATE exam_sessions SET flagged_count=flagged_count+1 WHERE id=?", (sid,))
        conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.route("/api/security-violation", methods=["POST"])
@student_session_required
def security_violation():
    require_csrf()
    sid = session["student_session_id"]
    data = request.get_json(silent=True) or {}
    source = str(data.get("source", "focus_escape"))[:80]
    allowed_sources = {"fullscreen_exit", "tab_hidden", "window_blur"}
    if source not in allowed_sources:
        source = "focus_escape"

    conn = connect()
    ex = conn.execute("SELECT * FROM exam_sessions WHERE id=?", (sid,)).fetchone()
    if not ex or ex["status"] != "in_progress":
        conn.close()
        return jsonify({"ok": False, "error": "Exam is not active"}), 409
    if session_security_mode(conn, ex) != "strict":
        clear_non_strict_security(conn, sid)
        conn.commit()
        ex = conn.execute("SELECT * FROM exam_sessions WHERE id=?", (sid,)).fetchone()
        state = get_security_state(ex)
        state.update({"permanent": False, "pending": False, "resume_required": False, "temp_remaining": 0})
        conn.close()
        return jsonify({"ok": True, "ignored": True, **state})

    existing_state = get_security_state(ex)
    if existing_state["permanent"] or existing_state["pending"] or existing_state["resume_required"] or existing_state["temp_remaining"] > 0:
        conn.close()
        return jsonify({"ok": True, "already_locked": True, **existing_state})

    # A single Alt+Tab can emit blur + visibility + fullscreen events. Treat
    # events arriving within a few seconds as one integrity violation.
    last = conn.execute(
        "SELECT created_at FROM proctor_events WHERE session_id=? AND event_type='security_violation' ORDER BY id DESC LIMIT 1",
        (sid,),
    ).fetchone()
    now = datetime.now(APP_TZ)
    if last:
        last_at = parse_iso(last["created_at"])
        if last_at and (now - last_at).total_seconds() < SECURITY_DEDUPE_SECONDS:
            state = get_security_state(ex)
            conn.close()
            return jsonify({"ok": True, "deduped": True, **state})

    current_count = int(ex["violation_count"] or 0) + 1
    created_at = iso_now()
    permanent = current_count >= MAX_SECURITY_VIOLATIONS
    if permanent:
        conn.execute(
            "UPDATE exam_sessions SET violation_count=?, security_locked=1, temp_locked_until=NULL, pending_blackout=0, security_resume_required=0, flagged_count=flagged_count+1 WHERE id=?",
            (current_count, sid),
        )
        detail = f"Violation {current_count}: {source}. Attempt permanently locked pending instructor unlock."
    else:
        # Queue the blackout until the exam regains focus. A browser cannot
        # black out another application while Alt+Tab is active, so serving
        # the full 15-second penalty on return is the strongest reliable web behavior.
        conn.execute(
            "UPDATE exam_sessions SET violation_count=?, pending_blackout=1, temp_locked_until=NULL, flagged_count=flagged_count+1 WHERE id=?",
            (current_count, sid),
        )
        detail = f"Violation {current_count}: {source}. {TEMP_LOCK_SECONDS}-second blackout queued for the student's return to the exam."
    conn.execute(
        "INSERT INTO proctor_events(session_id,event_type,detail,created_at) VALUES (?,?,?,?)",
        (sid, "security_violation", detail, created_at),
    )
    conn.commit()
    updated = conn.execute("SELECT * FROM exam_sessions WHERE id=?", (sid,)).fetchone()
    state = get_security_state(updated)
    conn.close()
    return jsonify({"ok": True, **state})


@app.route("/api/security-start-lock", methods=["POST"])
@student_session_required
def security_start_lock():
    require_csrf()
    sid = session["student_session_id"]
    conn = connect()
    ex = conn.execute("SELECT * FROM exam_sessions WHERE id=?", (sid,)).fetchone()
    if not ex or ex["status"] != "in_progress":
        conn.close()
        return jsonify({"ok": False, "error": "Exam is not active"}), 409
    if session_security_mode(conn, ex) != "strict":
        clear_non_strict_security(conn, sid)
        conn.commit()
        ex = conn.execute("SELECT * FROM exam_sessions WHERE id=?", (sid,)).fetchone()
        state = get_security_state(ex)
        state.update({"permanent": False, "pending": False, "resume_required": False, "temp_remaining": 0})
        conn.close()
        return jsonify({"ok": True, "ignored": True, **state})
    state = get_security_state(ex)
    if state["permanent"] or state["resume_required"] or state["temp_remaining"] > 0:
        conn.close()
        return jsonify({"ok": True, **state})
    if state["pending"]:
        lock_until = datetime.now(APP_TZ) + timedelta(seconds=TEMP_LOCK_SECONDS)
        conn.execute(
            "UPDATE exam_sessions SET pending_blackout=0, temp_locked_until=? WHERE id=?",
            (lock_until.isoformat(timespec="seconds"), sid),
        )
        conn.execute(
            "INSERT INTO proctor_events(session_id,event_type,detail,created_at) VALUES (?,?,?,?)",
            (sid, "security_blackout_started", f"{TEMP_LOCK_SECONDS}-second blackout penalty started after returning to the exam.", iso_now()),
        )
        conn.commit()
        ex = conn.execute("SELECT * FROM exam_sessions WHERE id=?", (sid,)).fetchone()
        state = get_security_state(ex)
    conn.close()
    return jsonify({"ok": True, **state})


@app.route("/api/security-resume", methods=["POST"])
@student_session_required
def security_resume():
    """Complete an instructor-granted unlock only after secure mode is restored.

    The unlock action deliberately leaves a server-side resume gate in place.
    Answers, navigation persistence, review flags, and submission remain blocked
    until the student's browser explicitly re-enters the secure display flow.
    """
    require_csrf()
    sid = session["student_session_id"]
    data = request.get_json(silent=True) or {}
    conn = connect()
    ex = conn.execute("SELECT * FROM exam_sessions WHERE id=?", (sid,)).fetchone()
    if not ex or ex["status"] != "in_progress":
        conn.close()
        return jsonify({"ok": False, "error": "Exam is not active"}), 409
    if session_security_mode(conn, ex) != "strict":
        clear_non_strict_security(conn, sid)
        conn.commit()
        ex = conn.execute("SELECT * FROM exam_sessions WHERE id=?", (sid,)).fetchone()
        state = get_security_state(ex)
        state.update({"permanent": False, "pending": False, "resume_required": False, "temp_remaining": 0})
        conn.close()
        return jsonify({"ok": True, "ignored": True, **state})
    if not bool(data.get("secure_active")):
        conn.close()
        return jsonify({"ok": False, "error": "secure_mode_required"}), 409
    state = get_security_state(ex)
    if state["permanent"] or state["pending"] or state["temp_remaining"] > 0:
        conn.close()
        return jsonify({"ok": False, "locked": True, **state}), 423
    if state["resume_required"]:
        conn.execute("UPDATE exam_sessions SET security_resume_required=0 WHERE id=?", (sid,))
        conn.execute(
            "INSERT INTO proctor_events(session_id,event_type,detail,created_at) VALUES (?,?,?,?)",
            (sid, "security_resume_confirmed", "Student re-entered secure display mode after instructor unlock.", iso_now()),
        )
        conn.commit()
        ex = conn.execute("SELECT * FROM exam_sessions WHERE id=?", (sid,)).fetchone()
        state = get_security_state(ex)
    conn.close()
    return jsonify({"ok": True, **state})


@app.route("/api/security-status")
@student_session_required
def security_status():
    sid = session["student_session_id"]
    conn = connect()
    ex = conn.execute("SELECT * FROM exam_sessions WHERE id=?", (sid,)).fetchone()
    if not ex:
        conn.close()
        return jsonify({"ok": False}), 404
    if session_security_mode(conn, ex) != "strict":
        clear_non_strict_security(conn, sid)
        conn.commit()
        ex = conn.execute("SELECT * FROM exam_sessions WHERE id=?", (sid,)).fetchone()
        state = get_security_state(ex)
        state.update({"permanent": False, "pending": False, "resume_required": False, "temp_remaining": 0})
        conn.close()
        return jsonify({"ok": True, "ignored": True, **state})
    state = get_security_state(ex)
    conn.close()
    return jsonify({"ok": True, **state})


@app.route("/api/exam-tools", methods=["GET", "POST"])
@student_session_required
def exam_tools():
    """Persistent CSDC101 calculator/scratchpad workspace state.

    The calculator itself is client-side; only the typed notes and normalized
    whiteboard strokes are stored so a refresh does not erase working paper.
    """
    sid = session["student_session_id"]
    conn = connect()
    ex = conn.execute("SELECT * FROM exam_sessions WHERE id=?", (sid,)).fetchone()
    if not ex:
        conn.close()
        return jsonify({"ok": False, "error": "session_not_found"}), 404
    if session_subject_code(conn, ex) != "CSDC101":
        conn.close()
        return jsonify({"ok": False, "error": "tools_not_available"}), 403
    if request.method == "GET":
        board = []
        raw = ex["scratch_board_json"] if "scratch_board_json" in ex.keys() else None
        if raw:
            try:
                board = json.loads(raw)
            except Exception:
                board = []
        payload = {
            "ok": True,
            "note": (ex["scratch_note"] or "") if "scratch_note" in ex.keys() else "",
            "board": board if isinstance(board, list) else [],
        }
        conn.close()
        return jsonify(payload)

    require_csrf()
    if ex["status"] != "in_progress":
        conn.close()
        return jsonify({"ok": False, "error": "Exam is not active"}), 409
    data = request.get_json(silent=True) or {}
    note = str(data.get("note", ""))[:8000]
    board = data.get("board", [])
    if not isinstance(board, list):
        board = []
    board = board[:1500]
    try:
        board_json = json.dumps(board, separators=(",", ":"))
    except (TypeError, ValueError):
        conn.close()
        return jsonify({"ok": False, "error": "invalid_board"}), 400
    if len(board_json) > 120000:
        conn.close()
        return jsonify({"ok": False, "error": "board_too_large"}), 413
    conn.execute(
        "UPDATE exam_sessions SET scratch_note=?, scratch_board_json=? WHERE id=?",
        (note, board_json, sid),
    )
    conn.commit()
    conn.close()
    return jsonify({"ok": True})




def _assessment_resource_allowed(assessment_id):
    if session.get("admin_id"):
        return True
    sid = session.get("student_session_id")
    if not sid:
        return False
    conn = connect()
    try:
        row = conn.execute(
            """SELECT COALESCE(e.assessment_id,b.assessment_id) AS assessment_id
               FROM exam_sessions e JOIN batches b ON b.id=e.batch_id WHERE e.id=?""",
            (sid,),
        ).fetchone()
        return bool(row and int(row["assessment_id"] or 0) == int(assessment_id))
    finally:
        conn.close()


@app.get("/assessment/<int:assessment_id>/workbench/notebook")
def assessment_workbench_notebook(assessment_id):
    if not _assessment_resource_allowed(assessment_id):
        abort(403)
    conn = connect()
    try:
        row = conn.execute("SELECT filename,notebook_json FROM assessment_notebooks WHERE assessment_id=?", (assessment_id,)).fetchone()
        if not row:
            abort(404)
        data = row["notebook_json"].encode("utf-8")
        return send_file(io.BytesIO(data), mimetype="application/x-ipynb+json", as_attachment=True, download_name=row["filename"])
    finally:
        conn.close()


@app.get("/assessment/<int:assessment_id>/workbench/assets/<path:asset_path>")
def assessment_workbench_asset(assessment_id, asset_path):
    if not _assessment_resource_allowed(assessment_id):
        abort(403)
    import base64
    safe = str(asset_path or "").replace("\\", "/").lstrip("/")
    if not safe or ".." in safe.split("/"):
        abort(404)
    conn = connect()
    try:
        row = conn.execute(
            "SELECT mime_type,content_b64 FROM assessment_notebook_assets WHERE assessment_id=? AND asset_path=?",
            (assessment_id, safe),
        ).fetchone()
        if not row:
            abort(404)
        try:
            data = base64.b64decode(row["content_b64"])
        except Exception:
            abort(404)
        return Response(data, mimetype=row["mime_type"] or "application/octet-stream")
    finally:
        conn.close()


@app.get("/assessment/<int:assessment_id>/workbench/package")
def assessment_workbench_package(assessment_id):
    if not _assessment_resource_allowed(assessment_id):
        abort(403)
    import base64
    import zipfile
    conn = connect()
    try:
        nb = conn.execute("SELECT filename,notebook_json,package_name FROM assessment_notebooks WHERE assessment_id=?", (assessment_id,)).fetchone()
        if not nb:
            abort(404)
        assets = conn.execute(
            "SELECT asset_path,content_b64 FROM assessment_notebook_assets WHERE assessment_id=? ORDER BY asset_path", (assessment_id,)
        ).fetchall()
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.writestr(nb["filename"], nb["notebook_json"])
            for asset in assets:
                zf.writestr(asset["asset_path"], base64.b64decode(asset["content_b64"]))
        out.seek(0)
        name = nb["package_name"] or f"assessment-{assessment_id}-workbench.zip"
        if not str(name).lower().endswith(".zip"):
            name = f"{name}.zip"
        return send_file(out, mimetype="application/zip", as_attachment=True, download_name=name)
    finally:
        conn.close()


@app.route("/api/lab-workbench", methods=["GET", "POST"])
@student_session_required
def lab_workbench_state():
    """Persist editable CSEC303 notebook cell source across days and devices.

    The Pyodide process itself is intentionally ephemeral; Python variables and
    generated figures are reconstructed by rerunning cells. The student's edited
    cell source is server-side assessment state and survives sign-out/reconnect.
    """
    sid = session["student_session_id"]
    conn = connect()
    ex = conn.execute("SELECT * FROM exam_sessions WHERE id=?", (sid,)).fetchone()
    if not ex:
        conn.close(); return jsonify({"ok": False, "error": "session_not_found"}), 404
    subject = session_subject_code(conn, ex)
    if subject != "CSEC303":
        conn.close(); return jsonify({"ok": False, "error": "not_available"}), 404
    if request.method == "GET":
        raw = ex["lab_workbench_json"] if "lab_workbench_json" in ex.keys() else None
        conn.close()
        try: state = json.loads(raw) if raw else {}
        except (TypeError, ValueError): state = {}
        return jsonify({"ok": True, "cells": state if isinstance(state, dict) else {}})
    require_csrf()
    if ex["status"] != "in_progress":
        conn.close(); return jsonify({"ok": False, "error": "submitted"}), 409
    data = request.get_json(silent=True) or {}
    cells = data.get("cells") or {}
    if not isinstance(cells, dict):
        conn.close(); return jsonify({"ok": False, "error": "invalid_state"}), 400
    cleaned = {}
    total = 0
    for key, value in list(cells.items())[:100]:
        key = str(key)[:80]
        value = str(value)
        total += len(value)
        if total > 250000:
            conn.close(); return jsonify({"ok": False, "error": "state_too_large"}), 413
        cleaned[key] = value
    conn.execute("UPDATE exam_sessions SET lab_workbench_json=? WHERE id=?", (json.dumps(cleaned), sid))
    conn.commit(); conn.close()
    return jsonify({"ok": True, "saved": len(cleaned)})


@app.route("/api/chat/messages")
@student_session_required
def student_chat_messages():
    sid = session["student_session_id"]
    conn = connect()
    ex = conn.execute("SELECT status FROM exam_sessions WHERE id=?", (sid,)).fetchone()
    if not ex:
        conn.close()
        return jsonify({"ok": False}), 404
    rows = conn.execute(
        "SELECT id,sender,message,created_at FROM exam_messages WHERE session_id=? ORDER BY id",
        (sid,),
    ).fetchall()
    if request.args.get("mark_read") == "1":
        conn.execute(
            "UPDATE exam_messages SET read_at=? WHERE session_id=? AND sender='instructor' AND read_at IS NULL",
            (iso_now(), sid),
        )
        conn.commit()
    payload = [dict(r) for r in rows]
    conn.close()
    return jsonify({"ok": True, "messages": payload})


@app.route("/api/chat/send", methods=["POST"])
@student_session_required
def student_chat_send():
    require_csrf()
    sid = session["student_session_id"]
    data = request.get_json(silent=True) or {}
    message = re.sub(r"\s+", " ", str(data.get("message", "")).strip())[:1000]
    if not message:
        return jsonify({"ok": False, "error": "Message is empty"}), 400
    conn = connect()
    ex = conn.execute("SELECT status FROM exam_sessions WHERE id=?", (sid,)).fetchone()
    if not ex or ex["status"] != "in_progress":
        conn.close()
        return jsonify({"ok": False, "error": "Exam is not active"}), 409
    conn.execute(
        "INSERT INTO exam_messages(session_id,sender,message,created_at) VALUES (?,?,?,?)",
        (sid, "student", message, iso_now()),
    )
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.route("/submit", methods=["POST"])
@student_session_required
def submit_exam():
    require_csrf()
    sid = session["student_session_id"]
    conn = connect()
    ex = conn.execute("SELECT * FROM exam_sessions WHERE id=?", (sid,)).fetchone()
    if ex:
        batch = conn.execute("SELECT * FROM batches WHERE id=?", (ex["batch_id"],)).fetchone()
        remaining, _ = session_remaining(ex, batch)
        if remaining is not None and remaining <= 0:
            finalize_exam(conn, sid, "time_limit_reached")
            conn.close()
            return redirect(url_for("result"))
        blocked, security = session_security_blocked(ex, conn)
        if blocked:
            conn.close()
            flash("This attempt is security-locked. Ask the instructor to unlock it before submitting.", "error")
            return redirect(url_for("exam"))
    finalize_exam(conn, sid, "student_submitted")
    conn.close()
    return redirect(url_for("result"))


@app.route("/result")
@student_session_required
def result():
    sid = session["student_session_id"]
    conn = connect()
    ex = conn.execute(
        """SELECT e.*, b.name AS batch_name, b.reveal_score, COALESCE(b.reveal_answers,0) AS reveal_answers, b.assessment_type,
                  a.title AS assessment_title, a.display_type, a.assessment_type AS configured_type,
                  COALESCE(a.reveal_answers,0) AS assessment_reveal_answers, s.code AS subject_code
           FROM exam_sessions e JOIN batches b ON b.id=e.batch_id
           LEFT JOIN assessments a ON a.id=e.assessment_id
           LEFT JOIN subjects s ON s.id=a.subject_id
           WHERE e.id=?""", (sid,)
    ).fetchone()
    if not ex or ex["status"] != "submitted":
        conn.close()
        return redirect(url_for("exam"))
    is_custom = (ex["configured_type"] == "custom") if "configured_type" in ex.keys() else False
    if is_custom:
        max_row = conn.execute(
            """SELECT COALESCE(SUM(q.points),0) AS max_score
               FROM session_questions sq JOIN questions q ON q.id=sq.question_id WHERE sq.session_id=?""",
            (sid,),
        ).fetchone()
        max_total = int(max_row["max_score"] or 0)
        display_total = ex["admin_total"] if ex["admin_total"] is not None else ex["auto_total"]
        display_p1 = ex["admin_part1_score"] if ex["admin_part1_score"] is not None else ex["part1_score"]
        display_p2 = 0
        display_bonus = 0
        is_midterm = False
    else:
        display_p1 = ex["admin_part1_score"] if ex["admin_part1_score"] is not None else ex["part1_score"]
        display_p2 = ex["admin_part2_score"] if ex["admin_part2_score"] is not None else ex["part2_score"]
        display_bonus = ex["admin_bonus_score"] if ex["admin_bonus_score"] is not None else ex["bonus_score"]
        is_midterm = ex["assessment_type"] == "midterm"
        max_total = 55 if is_midterm else 50
        computed_total = display_p1 + display_p2 + (display_bonus if is_midterm else 0)
        display_total = ex["admin_total"] if ex["admin_total"] is not None else computed_total
    can_review_answers = bool(ex["is_test"] or ex["reveal_answers"] or ex["assessment_reveal_answers"])
    review_items = []
    bonus_review = []
    if can_review_answers:
        qrows = conn.execute(
            """SELECT sq.q_order,sq.option_order,sq.selected_option,q.id AS question_id,q.prompt,q.code,q.explanation,q.points,
                      q.option_a,q.option_b,q.option_c,q.option_d,q.correct_option
               FROM session_questions sq JOIN questions q ON q.id=sq.question_id
               WHERE sq.session_id=? ORDER BY sq.q_order""", (sid,)
        ).fetchall()
        for q in qrows:
            option_map={"A":q["option_a"],"B":q["option_b"],"C":q["option_c"],"D":q["option_d"]}
            try: order=json.loads(q["option_order"] or "[]")
            except (TypeError,ValueError): order=["A","B","C","D"]
            if len(order)!=4: order=["A","B","C","D"]
            display_for_origin={origin:letter for letter,origin in zip("ABCD",order)}
            selected=q["selected_option"]
            correct=q["correct_option"]
            review_items.append({
                "order":q["q_order"],"prompt":q["prompt"],"code":q["code"],"points":q["points"],
                "selected_origin":selected,"correct_origin":correct,
                "selected_letter":display_for_origin.get(selected,"—"),"correct_letter":display_for_origin.get(correct,"—"),
                "selected_text":option_map.get(selected,"No answer"),"correct_text":option_map.get(correct,""),
                "is_correct":bool(selected and selected==correct),"explanation":q["explanation"] or "",
                "display_options":[{"letter":letter,"origin":origin,"text":option_map[origin]} for letter,origin in zip("ABCD",order)],
            })
        if ex["assessment_type"] == "midterm":
            brows=conn.execute(
                """SELECT sba.q_order,sba.answer_text,b.prompt,b.accepted_answer,b.topic
                   FROM session_bonus_answers sba JOIN bonus_questions b ON b.id=sba.bonus_question_id
                   WHERE sba.session_id=? ORDER BY sba.q_order""",(sid,)
            ).fetchall()
            bonus_review=[dict(x) for x in brows]
    conn.close()
    return render_template(
        "result.html", ex=ex, display_p1=display_p1, display_p2=display_p2,
        display_bonus=display_bonus, display_total=display_total, max_total=max_total,
        is_midterm=is_midterm, is_custom=is_custom, can_review_answers=can_review_answers,
        review_items=review_items, bonus_review=bonus_review, subject_code=(ex["subject_code"] or "CSDC101")
    )


@app.route("/logout")
def student_logout():
    for key in ("student_session_id", "pending_email", "pending_first_name", "pending_last_name", "pending_student_name", "pending_program", "pending_class_section", "pending_batch_id", "pending_session_key", "pending_google_sub", "google_student", "student_identity", "selected_subject_code"):
        session.pop(key, None)
    return redirect(url_for("student_login"))


@app.route("/admin/login", methods=["GET", "POST"])
def admin_login():
    # Student and Instructor View are intentionally separate entry points.
    # Classroom-enabled deployments render the teacher Google flow here; the
    # ?password=1 path remains the owner emergency sign-in.
    if request.method == "POST":
        require_csrf()
        guard_key = _admin_login_guard_key()
        retry_after = _admin_login_retry_after(guard_key)
        if retry_after:
            flash(f"Too many failed sign-in attempts. Try again in about {max(1, retry_after // 60)} minute(s).", "error")
            response = app.make_response(render_template("admin_login.html"))
            response.status_code = 429
            response.headers["Retry-After"] = str(retry_after)
            return response
        username = request.form.get("username", "")
        password = request.form.get("password", "")
        conn = connect()
        admin = conn.execute("SELECT * FROM admins WHERE username=?", (username,)).fetchone()
        conn.close()
        if admin and instructor_scope.active() and admin["role"] != "owner":
            admin = None  # Classroom mode: password login is an owner-only emergency path
        if admin and ("active" not in admin.keys() or admin["active"]) and check_password_hash(admin["password_hash"], password):
            session.clear()
            session["admin_id"] = admin["id"]
            session["admin_role"] = admin["role"] if "role" in admin.keys() else "owner"
            session["admin_display_name"] = (admin["display_name"] if "display_name" in admin.keys() else None) or admin["username"]
            session["admin_auth_method"] = "password"
            session["admin_authenticated"] = True
            session["admin_authenticated_at"] = iso_now()
            csrf_token()
            _clear_admin_login_failures(guard_key)
            return redirect(url_for("nextgen.workspace"))
        _record_admin_login_failure(guard_key)
        flash("Invalid administrator credentials.", "error")
    return render_template("admin_login.html")


@app.route("/admin/logout")
def admin_logout():
    session.clear()
    return redirect(url_for("admin_login"))


@app.route("/admin")
@admin_required
def admin_dashboard():
    requested = request.args.get("assessment", "").strip().lower()
    if not requested:
        conn = connect()
        assessments = conn.execute(
            """SELECT a.*, s.code AS subject_code, s.name AS subject_name,
                      CASE WHEN a.assessment_type='dryrun' THEN COALESCE(a.access_code,(SELECT b.access_code FROM batches b WHERE b.assessment_id=a.id ORDER BY b.id LIMIT 1)) ELSE a.access_code END AS effective_access_code,
                      (SELECT b.id FROM batches b WHERE b.assessment_id=a.id ORDER BY b.id LIMIT 1) AS delivery_batch_id,
                      (SELECT COUNT(*) FROM questions q WHERE q.assessment_id=a.id AND COALESCE(q.active,1)=1) AS question_count,
                      (SELECT COALESCE(SUM(q.points),0) FROM questions q WHERE q.assessment_id=a.id AND COALESCE(q.active,1)=1) AS max_score,
                      (SELECT COUNT(*) FROM exam_sessions e WHERE e.assessment_id=a.id AND COALESCE(e.is_test,0)=0) AS attempt_count,
                      (SELECT COUNT(*) FROM exam_sessions e WHERE e.assessment_id=a.id AND COALESCE(e.is_test,0)=0 AND e.status='submitted') AS submitted_count,
                      (SELECT AVG(e.auto_total) FROM exam_sessions e WHERE e.assessment_id=a.id AND COALESCE(e.is_test,0)=0 AND e.status='submitted') AS avg_score,
                      (SELECT COUNT(*) FROM coding_sessions cs JOIN programming_labs pl ON pl.id=cs.lab_id WHERE pl.assessment_id=a.id) AS coding_attempt_count
               FROM assessments a JOIN subjects s ON s.id=a.subject_id
               WHERE a.deleted_at IS NULL
               ORDER BY a.active DESC, a.created_at DESC, a.id DESC"""
        ).fetchall()
        subjects = conn.execute("SELECT * FROM subjects WHERE active=1 ORDER BY code,name").fetchall()
        totals = conn.execute(
            """SELECT COUNT(*) AS assessment_count,
                      SUM(CASE WHEN active=1 THEN 1 ELSE 0 END) AS active_count
               FROM assessments WHERE deleted_at IS NULL"""
        ).fetchone()
        active_sessions = conn.execute(
            "SELECT COUNT(*) AS c FROM exam_sessions WHERE status='in_progress' AND COALESCE(is_test,0)=0"
        ).fetchone()["c"]
        if instructor_scope.list_scoped():
            # Classroom mode: only the signed-in teacher's own courses (owners too, unless "Show all").
            assessments = instructor_scope.filter_rows(conn, assessments)
            subjects = instructor_scope.filter_rows(conn, subjects, kind="subject")
            totals = {"assessment_count": len(assessments),
                      "active_count": sum(1 for a in assessments if a["active"])}
            visible = instructor_scope.visible_assessment_ids(conn)
            active_sessions = sum(1 for r in conn.execute(
                "SELECT assessment_id FROM exam_sessions WHERE status='in_progress' AND COALESCE(is_test,0)=0"
            ).fetchall() if r["assessment_id"] in visible)
        conn.close()
        return render_template(
            "admin_assessments_dashboard.html", assessments=assessments, subjects=subjects,
            totals=totals, active_sessions=active_sessions, assessment=""
        )

    assessment = requested
    if assessment not in {"midterm", "posttest"}:
        assessment = "posttest"
    session["admin_assessment"] = assessment
    conn = connect()
    batches = conn.execute(
        """SELECT b.*,
           COUNT(e.id) AS session_count,
           SUM(CASE WHEN e.status='submitted' THEN 1 ELSE 0 END) AS submitted_count
           FROM batches b LEFT JOIN exam_sessions e ON e.batch_id=b.id AND COALESCE(e.is_test,0)=0
           WHERE b.assessment_type=?
           GROUP BY b.id ORDER BY b.slot""",
        (assessment,),
    ).fetchall()
    sessions = conn.execute(
        """SELECT e.*, b.name AS batch_name,
                  (SELECT COUNT(*) FROM exam_messages m WHERE m.session_id=e.id AND m.sender='student' AND m.read_at IS NULL) AS unread_messages
           FROM exam_sessions e JOIN batches b ON b.id=e.batch_id
           WHERE COALESCE(e.is_test,0)=0 AND b.assessment_type=?
           ORDER BY e.id DESC LIMIT 100""",
        (assessment,),
    ).fetchall()
    qcounts = conn.execute(
        """SELECT q.part, COUNT(*) c FROM questions q
           JOIN batches b ON b.slot=q.batch_slot
           WHERE b.assessment_type=? AND COALESCE(q.active,1)=1 GROUP BY q.part ORDER BY q.part""",
        (assessment,),
    ).fetchall()
    summary = conn.execute(
        """SELECT COUNT(*) AS total_sessions,
                  SUM(CASE WHEN e.status='submitted' THEN 1 ELSE 0 END) AS submitted,
                  AVG(CASE WHEN e.status='submitted' THEN e.auto_total END) AS avg_score
           FROM exam_sessions e JOIN batches b ON b.id=e.batch_id
           WHERE COALESCE(e.is_test,0)=0 AND b.assessment_type=?""",
        (assessment,),
    ).fetchone()
    conn.close()
    counts = {r["part"]: r["c"] for r in qcounts}
    return render_template("admin_dashboard.html", batches=batches, sessions=sessions, counts=counts, summary=summary, assessment=assessment, subject_code="CSDC101")


@app.route("/admin/assessment/<int:assessment_id>/export-scores")
@admin_required
def admin_export_scores(assessment_id):
    """Export real student scores for one assessment as an Excel workbook.

    Objective assessments use one worksheet per delivery set/batch. Programming
    labs use one worksheet with task-level score columns. Instructor test/preview
    sessions are excluded from the export.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("Score export requires openpyxl. Run: pip install -r requirements.txt") from exc

    conn = connect()
    assessment_row = conn.execute(
        """SELECT a.*, s.code AS subject_code, s.name AS subject_name
           FROM assessments a JOIN subjects s ON s.id=a.subject_id
           WHERE a.id=?""",
        (assessment_id,),
    ).fetchone()
    if not assessment_row:
        conn.close()
        abort(404)

    wb = Workbook()
    wb.remove(wb.active)
    used_sheet_names = set()

    def clean_sheet_name(value, fallback="Scores"):
        name = re.sub(r"[\[\]:*?/\\]", "-", str(value or fallback)).strip() or fallback
        name = name[:31]
        base = name
        counter = 2
        while name.casefold() in used_sheet_names:
            suffix = f" ({counter})"
            name = f"{base[:31-len(suffix)]}{suffix}"
            counter += 1
        used_sheet_names.add(name.casefold())
        return name

    def write_sheet(title, headers, rows):
        ws = wb.create_sheet(clean_sheet_name(title))
        ws.append(headers)
        header_fill = PatternFill("solid", fgColor="EAF2FF")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="17365D")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for idx, header in enumerate(headers, start=1):
            max_len = len(str(header))
            for cell in ws[get_column_letter(idx)]:
                value = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, len(value))
            ws.column_dimensions[get_column_letter(idx)].width = min(max(max_len + 2, 11), 34)
        return ws

    atype = assessment_row["assessment_type"]
    if atype == "programming_lab":
        lab = conn.execute(
            "SELECT * FROM programming_labs WHERE assessment_id=? LIMIT 1",
            (assessment_id,),
        ).fetchone()
        tasks = []
        sessions = []
        if lab:
            tasks = conn.execute(
                "SELECT id,position,title,points FROM programming_tasks WHERE lab_id=? AND COALESCE(active,1)=1 ORDER BY position,id",
                (lab["id"],),
            ).fetchall()
            sessions = conn.execute(
                """SELECT * FROM coding_sessions
                   WHERE lab_id=? AND COALESCE(is_test,0)=0
                   ORDER BY COALESCE(last_name,student_name,email), COALESCE(first_name,''), email""",
                (lab["id"],),
            ).fetchall()
        headers = [
            "Student Name", "Email", "Program", "Section", "Status",
            "Started At", "Submitted At",
        ]
        headers.extend([f"Task {t['position']}: {t['title']} / {t['points']}" for t in tasks])
        headers.extend(["Total Score", "Maximum Score", "Flags", "Violations"])
        max_score = sum(float(t["points"] or 0) for t in tasks)
        rows = []
        for cs in sessions:
            progress = conn.execute(
                "SELECT task_id,score FROM coding_task_progress WHERE session_id=?",
                (cs["id"],),
            ).fetchall()
            scores = {int(r["task_id"]): float(r["score"] or 0) for r in progress}
            display_name = (
                f"{(cs['last_name'] or '').strip()}, {(cs['first_name'] or '').strip()}".strip(", ")
                if (cs["last_name"] or cs["first_name"]) else (cs["student_name"] or cs["email"])
            )
            row = [
                display_name, cs["email"], cs["program"] or "", cs["class_section"] or "",
                str(cs["status"] or "").replace("_", " ").title(), cs["started_at"] or "", cs["submitted_at"] or "",
            ]
            row.extend([scores.get(int(t["id"]), 0) for t in tasks])
            row.extend([float(cs["total_score"] or 0), max_score, int(cs["flagged_count"] or 0), int(cs["violation_count"] or 0)])
            rows.append(row)
        write_sheet("Programming Scores", headers, rows)
    else:
        batches = conn.execute(
            "SELECT * FROM batches WHERE assessment_id=? ORDER BY slot,id",
            (assessment_id,),
        ).fetchall()
        if not batches:
            # Keep export useful even for legacy/imported data with a missing batch link.
            batches = conn.execute(
                """SELECT DISTINCT b.* FROM batches b
                   JOIN exam_sessions e ON e.batch_id=b.id
                   WHERE e.assessment_id=? ORDER BY b.slot,b.id""",
                (assessment_id,),
            ).fetchall()

        for batch in batches or [None]:
            if batch is None:
                sessions = conn.execute(
                    """SELECT e.*,
                              (SELECT COALESCE(SUM(q.points),0) FROM session_questions sq JOIN questions q ON q.id=sq.question_id WHERE sq.session_id=e.id) AS assigned_max_score
                       FROM exam_sessions e
                       WHERE e.assessment_id=? AND COALESCE(e.is_test,0)=0
                       ORDER BY COALESCE(e.last_name,e.student_name,e.email), COALESCE(e.first_name,''), e.email""",
                    (assessment_id,),
                ).fetchall()
                sheet_title = "Scores"
                set_label = ""
            else:
                sessions = conn.execute(
                    """SELECT e.*,
                              (SELECT COALESCE(SUM(q.points),0) FROM session_questions sq JOIN questions q ON q.id=sq.question_id WHERE sq.session_id=e.id) AS assigned_max_score
                       FROM exam_sessions e
                       WHERE e.batch_id=? AND COALESCE(e.is_test,0)=0
                       ORDER BY COALESCE(e.last_name,e.student_name,e.email), COALESCE(e.first_name,''), e.email""",
                    (batch["id"],),
                ).fetchall()
                sheet_title = batch["name"] or f"Set {batch['slot']}"
                set_label = batch["batch_label"] or batch["name"] or ""

            if atype in {"midterm", "posttest"}:
                headers = [
                    "Student Name", "Email", "Program", "Section", "Set", "Status",
                    "Started At", "Submitted At", "Part I Score", "Part II Score",
                    "Bonus Score", "Automatic Total", "Final Score", "Maximum Score",
                    "Flags", "Violations",
                ]
            else:
                headers = [
                    "Student Name", "Email", "Program", "Section", "Set", "Status",
                    "Started At", "Submitted At", "Raw Correct", "Automatic Score",
                    "Final Score", "Maximum Score", "Flags", "Violations",
                ]
            rows = []
            for ex in sessions:
                display_name = (
                    f"{(ex['last_name'] or '').strip()}, {(ex['first_name'] or '').strip()}".strip(", ")
                    if (ex["last_name"] or ex["first_name"]) else (ex["student_name"] or ex["email"])
                )
                final_score = ex["admin_total"] if ex["admin_total"] is not None else ex["auto_total"]
                if atype == "midterm":
                    max_score = 55
                elif atype == "posttest":
                    max_score = 50
                else:
                    max_score = float(ex["assigned_max_score"] or 0)
                    if max_score.is_integer():
                        max_score = int(max_score)
                common = [
                    display_name, ex["email"], ex["program"] or "", ex["class_section"] or "", set_label,
                    str(ex["status"] or "").replace("_", " ").title(), ex["started_at"] or "", ex["submitted_at"] or "",
                ]
                if atype in {"midterm", "posttest"}:
                    p1 = ex["admin_part1_score"] if ex["admin_part1_score"] is not None else ex["part1_score"]
                    p2 = ex["admin_part2_score"] if ex["admin_part2_score"] is not None else ex["part2_score"]
                    bonus = ex["admin_bonus_score"] if ex["admin_bonus_score"] is not None else ex["bonus_score"]
                    row = common + [p1, p2, bonus, ex["auto_total"], final_score, max_score, int(ex["flagged_count"] or 0), int(ex["violation_count"] or 0)]
                else:
                    row = common + [int(ex["part1_correct"] or 0), ex["auto_total"], final_score, max_score, int(ex["flagged_count"] or 0), int(ex["violation_count"] or 0)]
                rows.append(row)
            write_sheet(sheet_title, headers, rows)

    conn.close()
    if not wb.worksheets:
        write_sheet("Scores", ["Student Name", "Email", "Score"], [])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename_base = re.sub(r"[^A-Za-z0-9._-]+", "_", assessment_row["title"] or "assessment").strip("._") or "assessment"
    return send_file(
        output,
        as_attachment=True,
        download_name=f"{filename_base}_scores.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/admin/assessment/<int:assessment_id>/export-scores.csv")
@admin_required
def admin_export_scores_csv(assessment_id):
    """Export real student scores as one portable UTF-8 CSV file.

    Multi-set assessments are flattened into one file with a Set / Delivery
    column. The existing XLSX export remains the richer option when separate
    worksheets per set are preferred.
    """
    conn = connect()
    assessment_row = conn.execute(
        """SELECT a.*, s.code AS subject_code, s.name AS subject_name
           FROM assessments a JOIN subjects s ON s.id=a.subject_id
           WHERE a.id=?""",
        (assessment_id,),
    ).fetchone()
    if not assessment_row:
        conn.close()
        abort(404)

    sio = io.StringIO(newline="")
    writer = csv.writer(sio)
    atype = assessment_row["assessment_type"]

    if atype == "programming_lab":
        lab = conn.execute("SELECT * FROM programming_labs WHERE assessment_id=? LIMIT 1", (assessment_id,)).fetchone()
        tasks = conn.execute(
            "SELECT id,position,title,points FROM programming_tasks WHERE lab_id=? AND COALESCE(active,1)=1 ORDER BY position,id",
            (lab["id"],),
        ).fetchall() if lab else []
        writer.writerow([
            "Student Name", "Email", "Program", "Section", "Status", "Started At", "Submitted At",
            *[f"Task {t['position']}: {t['title']} / {t['points']}" for t in tasks],
            "Total Score", "Maximum Score", "Flags", "Violations",
        ])
        max_score = sum(float(t["points"] or 0) for t in tasks)
        sessions = conn.execute(
            """SELECT * FROM coding_sessions WHERE lab_id=? AND COALESCE(is_test,0)=0
               ORDER BY COALESCE(last_name,student_name,email), COALESCE(first_name,''), email""",
            (lab["id"],),
        ).fetchall() if lab else []
        for cs in sessions:
            progress = conn.execute("SELECT task_id,score FROM coding_task_progress WHERE session_id=?", (cs["id"],)).fetchall()
            scores = {int(r["task_id"]): float(r["score"] or 0) for r in progress}
            display_name = (
                f"{(cs['last_name'] or '').strip()}, {(cs['first_name'] or '').strip()}".strip(", ")
                if (cs["last_name"] or cs["first_name"]) else (cs["student_name"] or cs["email"])
            )
            writer.writerow([
                display_name, cs["email"], cs["program"] or "", cs["class_section"] or "",
                str(cs["status"] or "").replace("_", " ").title(), cs["started_at"] or "", cs["submitted_at"] or "",
                *[scores.get(int(t["id"]), 0) for t in tasks],
                float(cs["total_score"] or 0), max_score, int(cs["flagged_count"] or 0), int(cs["violation_count"] or 0),
            ])
    else:
        if atype in {"midterm", "posttest"}:
            headers = [
                "Student Name", "Email", "Program", "Section", "Set / Delivery", "Status",
                "Started At", "Submitted At", "Part I Score", "Part II Score", "Bonus Score",
                "Automatic Total", "Final Score", "Maximum Score", "Flags", "Violations",
            ]
        else:
            headers = [
                "Student Name", "Email", "Program", "Section", "Set / Delivery", "Status",
                "Started At", "Submitted At", "Raw Correct", "Automatic Score", "Final Score",
                "Maximum Score", "Flags", "Violations",
            ]
        writer.writerow(headers)
        sessions = conn.execute(
            """SELECT e.*, b.name AS delivery_name, b.batch_label,
                      (SELECT COALESCE(SUM(q.points),0) FROM session_questions sq
                       JOIN questions q ON q.id=sq.question_id WHERE sq.session_id=e.id) AS assigned_max_score
               FROM exam_sessions e JOIN batches b ON b.id=e.batch_id
               WHERE COALESCE(e.is_test,0)=0 AND (e.assessment_id=? OR b.assessment_id=?)
               ORDER BY b.slot, COALESCE(e.last_name,e.student_name,e.email), COALESCE(e.first_name,''), e.email""",
            (assessment_id, assessment_id),
        ).fetchall()
        for ex in sessions:
            display_name = (
                f"{(ex['last_name'] or '').strip()}, {(ex['first_name'] or '').strip()}".strip(", ")
                if (ex["last_name"] or ex["first_name"]) else (ex["student_name"] or ex["email"])
            )
            set_label = ex["batch_label"] or ex["delivery_name"] or ""
            final_score = ex["admin_total"] if ex["admin_total"] is not None else ex["auto_total"]
            if atype == "midterm":
                max_score = 55
            elif atype == "posttest":
                max_score = 50
            else:
                max_score = float(ex["assigned_max_score"] or 0)
                if max_score.is_integer():
                    max_score = int(max_score)
            common = [
                display_name, ex["email"], ex["program"] or "", ex["class_section"] or "", set_label,
                str(ex["status"] or "").replace("_", " ").title(), ex["started_at"] or "", ex["submitted_at"] or "",
            ]
            if atype in {"midterm", "posttest"}:
                p1 = ex["admin_part1_score"] if ex["admin_part1_score"] is not None else ex["part1_score"]
                p2 = ex["admin_part2_score"] if ex["admin_part2_score"] is not None else ex["part2_score"]
                bonus = ex["admin_bonus_score"] if ex["admin_bonus_score"] is not None else ex["bonus_score"]
                writer.writerow(common + [p1, p2, bonus, ex["auto_total"], final_score, max_score, int(ex["flagged_count"] or 0), int(ex["violation_count"] or 0)])
            else:
                writer.writerow(common + [int(ex["part1_correct"] or 0), ex["auto_total"], final_score, max_score, int(ex["flagged_count"] or 0), int(ex["violation_count"] or 0)])

    conn.close()
    filename_base = re.sub(r"[^A-Za-z0-9._-]+", "_", assessment_row["title"] or "assessment").strip("._") or "assessment"
    data = chr(0xFEFF) + sio.getvalue()
    return Response(
        data,
        mimetype="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename_base}_scores.csv"'},
    )


@app.route("/admin/batch/<int:batch_id>", methods=["GET", "POST"])
@admin_required
def admin_batch(batch_id):
    conn = connect()
    batch = conn.execute("SELECT * FROM batches WHERE id=?", (batch_id,)).fetchone()
    if not batch:
        conn.close()
        abort(404)
    if request.method == "POST":
        require_csrf()
        open_at = request.form.get("open_at", "").strip() or None
        close_at = request.form.get("close_at", "").strip() or None
        duration = max(1, min(10080, int(request.form.get("duration_minutes", "90"))))
        reveal = 1 if request.form.get("reveal_score") == "on" else 0
        reveal_answers = 1 if request.form.get("reveal_answers") == "on" else 0
        active = 1 if request.form.get("active") == "on" else 0
        # HTML datetime-local is timezone-naive; store local wall time and parse it as server-local.
        try:
            conn.execute(
                """UPDATE batches SET open_at=?, close_at=?, duration_minutes=?, reveal_score=?, reveal_answers=?, active=?
                   WHERE id=?""",
                (open_at, close_at, duration, reveal, reveal_answers, active, batch_id),
            )
            conn.commit()
            flash("Batch settings updated.", "success")
        except Exception as exc:
            conn.rollback()
            flash(f"Could not update batch: {exc}", "error")
        batch = conn.execute(
            """SELECT b.*, COALESCE(s.code, sa.code, 'CSDC101') AS subject_code
               FROM batches b
               LEFT JOIN subjects s ON s.id=b.subject_id
               LEFT JOIN assessments a ON a.id=b.assessment_id
               LEFT JOIN subjects sa ON sa.id=a.subject_id
               WHERE b.id=?""",
            (batch_id,),
        ).fetchone()
    conn.close()
    return render_template("admin_batch.html", batch=batch, subject_code=batch["subject_code"])


@app.route("/admin/assessment/<assessment>/regenerate-keys", methods=["POST"])
@admin_required
def admin_regenerate_all_keys(assessment):
    require_csrf()
    assessment = assessment.strip().lower()
    if assessment not in {"midterm", "posttest"}:
        abort(404)
    conn = connect()
    batches = conn.execute(
        "SELECT * FROM batches WHERE assessment_type=? ORDER BY slot", (assessment,)
    ).fetchall()
    for batch in batches:
        new_key = unique_session_key(
            conn, batch["assessment_type"], batch["section"], batch["batch_label"]
        )
        conn.execute("UPDATE batches SET access_code=? WHERE id=?", (new_key, batch["id"]))
    conn.commit()
    conn.close()
    flash(f"Generated new session keys for all {len(batches)} sets. All previous keys are now invalid.", "success")
    return redirect(url_for("admin_dashboard", assessment=assessment))


@app.route("/admin/batch/<int:batch_id>/regenerate-key", methods=["POST"])
@admin_required
def admin_regenerate_key(batch_id):
    require_csrf()
    conn = connect()
    batch = conn.execute("SELECT * FROM batches WHERE id=?", (batch_id,)).fetchone()
    if not batch:
        conn.close()
        abort(404)
    new_key = unique_session_key(
        conn, batch["assessment_type"], batch["section"], batch["batch_label"]
    )
    conn.execute("UPDATE batches SET access_code=? WHERE id=?", (new_key, batch_id))
    conn.commit()
    conn.close()
    flash("A new session key was generated. The previous key is now invalid.", "success")
    return redirect(url_for("admin_batch", batch_id=batch_id))


@app.route("/admin/session/<int:sid>", methods=["GET", "POST"])
@admin_required
def admin_session_detail(sid):
    conn = connect()
    ex = conn.execute(
        """SELECT e.*, b.name AS batch_name, b.assessment_type,
                  a.title AS assessment_title, a.display_type, a.assessment_type AS configured_type,
                  s.code AS subject_code, s.name AS subject_name
           FROM exam_sessions e JOIN batches b ON b.id=e.batch_id
           LEFT JOIN assessments a ON a.id=e.assessment_id
           LEFT JOIN subjects s ON s.id=COALESCE(a.subject_id,b.subject_id) WHERE e.id=?""", (sid,)
    ).fetchone()
    if not ex:
        conn.close()
        abort(404)
    is_custom = (ex["configured_type"] == "custom") if "configured_type" in ex.keys() else False
    custom_max_score = 0
    if is_custom:
        max_row = conn.execute(
            """SELECT COALESCE(SUM(q.points),0) AS max_score FROM session_questions sq
               JOIN questions q ON q.id=sq.question_id WHERE sq.session_id=?""", (sid,)
        ).fetchone()
        custom_max_score = int(max_row["max_score"] or 0)
    if request.method == "POST":
        require_csrf()
        def optional_float(name):
            value = request.form.get(name, "").strip()
            return None if value == "" else float(value)
        try:
            p1 = optional_float("admin_part1_score")
            p2 = optional_float("admin_part2_score")
            bonus = optional_float("admin_bonus_score")
            total = optional_float("admin_total")
            note = request.form.get("admin_note", "").strip()[:1000]
            is_midterm = ex["assessment_type"] == "midterm"
            if not is_custom:
                if p1 is not None and not (0 <= p1 <= 30):
                    raise ValueError("Part I override must be from 0 to 30.")
                if p2 is not None and not (0 <= p2 <= 20):
                    raise ValueError("Part II override must be from 0 to 20.")
                if bonus is not None and not (0 <= bonus <= 5):
                    raise ValueError("Bonus override must be from 0 to 5.")
                max_total = 55 if is_midterm else 50
            else:
                p1 = p2 = bonus = None
                max_total = custom_max_score
            if total is not None and not (0 <= total <= max_total):
                raise ValueError(f"Total override must be from 0 to {max_total}.")
            conn.execute(
                """UPDATE exam_sessions SET admin_part1_score=?, admin_part2_score=?, admin_bonus_score=?, admin_total=?, admin_note=? WHERE id=?""",
                (p1, p2, bonus, total, note, sid),
            )
            conn.commit()
            flash("Score override saved.", "success")
        except ValueError as exc:
            flash(str(exc), "error")
        ex = conn.execute(
            """SELECT e.*, b.name AS batch_name, b.assessment_type,
                      a.title AS assessment_title, a.display_type, a.assessment_type AS configured_type
               FROM exam_sessions e JOIN batches b ON b.id=e.batch_id
               LEFT JOIN assessments a ON a.id=e.assessment_id WHERE e.id=?""", (sid,)
        ).fetchone()

    answers = conn.execute(
        """SELECT sq.q_order, sq.selected_option, q.part, q.topic, q.prompt, q.correct_option
           FROM session_questions sq JOIN questions q ON q.id=sq.question_id
           WHERE sq.session_id=? ORDER BY sq.q_order""", (sid,)
    ).fetchall()
    bonus_answers = conn.execute(
        """SELECT b.position, b.topic, b.prompt, b.accepted_answer, sba.answer_text
           FROM session_bonus_answers sba JOIN bonus_questions b ON b.id=sba.bonus_question_id
           WHERE sba.session_id=? ORDER BY sba.q_order""",
        (sid,),
    ).fetchall()
    bonus_audit = []
    for b in bonus_answers:
        row = dict(b)
        row["is_correct"] = short_answer_matches(b["answer_text"], b["accepted_answer"])
        bonus_audit.append(row)
    events = conn.execute(
        "SELECT * FROM proctor_events WHERE session_id=? ORDER BY id", (sid,)
    ).fetchall()
    messages = conn.execute(
        "SELECT * FROM exam_messages WHERE session_id=? ORDER BY id", (sid,)
    ).fetchall()
    conn.execute(
        "UPDATE exam_messages SET read_at=? WHERE session_id=? AND sender='student' AND read_at IS NULL",
        (iso_now(), sid),
    )
    conn.commit()
    security = get_security_state(ex)
    selected_assessment = None
    if ex["assessment_id"]:
        selected_assessment = conn.execute(
            """SELECT a.*,s.code AS subject_code,s.name AS subject_name
               FROM assessments a JOIN subjects s ON s.id=a.subject_id WHERE a.id=?""",
            (ex["assessment_id"],),
        ).fetchone()
    conn.close()
    return render_template("admin_session.html", ex=ex, answers=answers, bonus_audit=bonus_audit, events=events, messages=messages, security=security, is_custom=is_custom, custom_max_score=custom_max_score, selected_assessment=selected_assessment)


@app.route("/admin/session/<int:sid>/unlock", methods=["POST"])
@admin_required
def admin_unlock_session(sid):
    require_csrf()
    conn = connect()
    ex = conn.execute("SELECT * FROM exam_sessions WHERE id=?", (sid,)).fetchone()
    if not ex:
        conn.close()
        abort(404)
    conn.execute(
        "UPDATE exam_sessions SET security_locked=0, temp_locked_until=NULL, pending_blackout=0, security_resume_required=1 WHERE id=?",
        (sid,),
    )
    conn.execute(
        "INSERT INTO proctor_events(session_id,event_type,detail,created_at) VALUES (?,?,?,?)",
        (sid, "security_unlocked", f"Instructor granted another chance; violation count retained at {ex['violation_count'] or 0}. Secure-mode re-entry is required before answering can resume.", iso_now()),
    )
    conn.commit()
    conn.close()
    flash("Another chance was granted. The student must re-enter secure mode before answering or submitting; prior violations remain recorded.", "success")
    return redirect(url_for("admin_session_detail", sid=sid))


@app.route("/admin/session/<int:sid>/chat", methods=["POST"])
@admin_required
def admin_chat_send(sid):
    require_csrf()
    data = request.get_json(silent=True) if request.is_json else request.form
    data = data or {}
    message = re.sub(r"\s+", " ", str(data.get("message", "")).strip())[:1000]
    wants_json = request.is_json or "application/json" in request.headers.get("Accept", "")
    if not message:
        if wants_json:
            return jsonify({"ok": False, "error": "Enter a message before sending."}), 400
        flash("Enter a message before sending.", "error")
        return redirect(url_for("admin_session_detail", sid=sid))
    conn = connect()
    ex = conn.execute("SELECT id FROM exam_sessions WHERE id=?", (sid,)).fetchone()
    if not ex:
        conn.close()
        if wants_json:
            return jsonify({"ok": False, "error": "Session not found."}), 404
        abort(404)
    created_at = iso_now()
    cur = conn.execute(
        "INSERT INTO exam_messages(session_id,sender,message,created_at) VALUES (?,?,?,?) RETURNING id",
        (sid, "instructor", message, created_at),
    )
    message_id = cur.fetchone()[0]
    conn.commit(); conn.close()
    if wants_json:
        return jsonify({"ok": True, "message": {"id": message_id, "sender": "instructor", "message": message, "created_at": created_at}})
    return redirect(url_for("admin_session_detail", sid=sid))


@app.route("/admin/session/<int:sid>/chat/messages")
@admin_required
def admin_chat_messages(sid):
    conn = connect()
    if not conn.execute("SELECT 1 FROM exam_sessions WHERE id=?", (sid,)).fetchone():
        conn.close()
        return jsonify({"ok": False}), 404
    rows = conn.execute(
        "SELECT id,sender,message,created_at FROM exam_messages WHERE session_id=? ORDER BY id",
        (sid,),
    ).fetchall()
    conn.execute(
        "UPDATE exam_messages SET read_at=? WHERE session_id=? AND sender='student' AND read_at IS NULL",
        (iso_now(), sid),
    )
    conn.commit()
    payload = [dict(r) for r in rows]
    conn.close()
    return jsonify({"ok": True, "messages": payload})


def _message_assessment_context(conn):
    """Resolve the assessment context used by Messages.

    Messages follows the same assessment selected in Live Monitor. A direct
    session link takes precedence so clicking Message on a monitor card always
    opens that student's assessment, including CSEC303.
    """
    options = instructor_scope.filter_rows(conn, _admin_assessment_options(conn))
    if not options:
        return options, None

    selected = None
    raw_id = str(request.args.get("assessment_id") or "").strip()
    if raw_id.isdigit():
        selected = next((row for row in options if int(row["id"]) == int(raw_id)), None)

    if selected is None:
        raw_sid = str(request.args.get("session") or "").strip()
        if raw_sid.isdigit():
            row = conn.execute("""SELECT COALESCE(e.assessment_id,b.assessment_id) AS assessment_id FROM exam_sessions e JOIN batches b ON b.id=e.batch_id WHERE e.id=?""", (int(raw_sid),)).fetchone()
            if row and row["assessment_id"]:
                selected = next((item for item in options if int(item["id"]) == int(row["assessment_id"])), None)

    if selected is None:
        remembered = session.get("admin_assessment_id")
        if remembered:
            selected = next((row for row in options if int(row["id"]) == int(remembered)), None)

    if selected is None:
        # Prefer an assessment with a currently active attempt, then the normal
        # assessment ordering used elsewhere in Instructor View.
        active = conn.execute(
            """SELECT e.assessment_id, MAX(e.id) AS latest_session
               FROM exam_sessions e
               WHERE e.status='in_progress' AND COALESCE(e.is_test,0)=0
                 AND COALESCE(e.monitor_done,0)=0 AND e.assessment_id IS NOT NULL
               GROUP BY e.assessment_id ORDER BY latest_session DESC"""
        ).fetchall()
        active_ids = [int(row["assessment_id"]) for row in active]
        selected = next((item for aid in active_ids for item in options if int(item["id"]) == aid), None)

    if selected is None:
        selected = options[0]

    session["admin_assessment_id"] = int(selected["id"])
    if selected["assessment_type"] in {"midterm", "posttest"}:
        session["admin_assessment"] = selected["assessment_type"]
    return options, selected


def _message_threads(conn, assessment_id=None):
    params = []
    assessment_clause = ""
    if assessment_id:
        assessment_clause = " AND COALESCE(e.assessment_id,b.assessment_id)=?"
        params.append(int(assessment_id))

    rows = conn.execute(
        f"""SELECT e.id,e.email,e.first_name,e.last_name,e.student_name,e.program,e.class_section,e.status,
                  COALESCE(e.assessment_id,b.assessment_id) AS assessment_id,
                  a.title AS assessment_title,s.code AS subject_code,
                  b.name AS batch_name,b.assessment_type,
                  (SELECT COUNT(*) FROM exam_messages m WHERE m.session_id=e.id AND m.sender='student' AND m.read_at IS NULL) AS unread_messages,
                  (SELECT m.message FROM exam_messages m WHERE m.session_id=e.id ORDER BY m.id DESC LIMIT 1) AS last_message,
                  (SELECT m.created_at FROM exam_messages m WHERE m.session_id=e.id ORDER BY m.id DESC LIMIT 1) AS last_message_at
           FROM exam_sessions e
           JOIN batches b ON b.id=e.batch_id
           LEFT JOIN assessments a ON a.id=COALESCE(e.assessment_id,b.assessment_id)
           LEFT JOIN subjects s ON s.id=a.subject_id
           WHERE COALESCE(e.is_test,0)=0
             {assessment_clause}
             AND (
                  (e.status='in_progress' AND COALESCE(e.monitor_done,0)=0)
                  OR EXISTS(SELECT 1 FROM exam_messages mx WHERE mx.session_id=e.id)
             )
           ORDER BY CASE WHEN e.status='in_progress' AND COALESCE(e.monitor_done,0)=0 THEN 0 ELSE 1 END,
                    CASE WHEN (SELECT COUNT(*) FROM exam_messages mu WHERE mu.session_id=e.id AND mu.sender='student' AND mu.read_at IS NULL) > 0 THEN 0 ELSE 1 END,
                    COALESCE((SELECT m2.created_at FROM exam_messages m2 WHERE m2.session_id=e.id ORDER BY m2.id DESC LIMIT 1),e.started_at) DESC, e.id DESC""",
        tuple(params),
    ).fetchall()
    return [dict(r) for r in instructor_scope.filter_rows(conn, rows, key="assessment_id")]


@app.get("/admin/messages")
@admin_required
def admin_messages():
    conn = connect()
    assessment_options, selected_assessment = _message_assessment_context(conn)
    selected_id = int(selected_assessment["id"]) if selected_assessment else None
    threads = _message_threads(conn, selected_id)
    conn.close()
    return render_template(
        "admin_messages.html",
        threads=threads,
        assessment=(selected_assessment["assessment_type"] if selected_assessment else session.get("admin_assessment", "posttest")),
        assessment_options=assessment_options,
        selected_assessment=selected_assessment,
        page_subject_code=(selected_assessment["subject_code"] if selected_assessment else None),
    )


@app.get("/admin/messages/threads")
@admin_required
def admin_message_threads():
    conn = connect()
    assessment_options, selected_assessment = _message_assessment_context(conn)
    selected_id = int(selected_assessment["id"]) if selected_assessment else None
    threads = _message_threads(conn, selected_id)
    conn.close()
    return jsonify({
        "ok": True,
        "threads": threads,
        "assessment_id": selected_id,
        "assessment_title": selected_assessment["title"] if selected_assessment else None,
        "subject_code": selected_assessment["subject_code"] if selected_assessment else None,
    })


@app.get("/admin/messages/<int:sid>")
@admin_required
def admin_message_thread(sid):
    conn = connect()
    ex = conn.execute(
        """SELECT e.id,e.email,e.first_name,e.last_name,e.student_name,e.program,e.class_section,e.status,
                  COALESCE(e.assessment_id,b.assessment_id) AS assessment_id,
                  a.title AS assessment_title,s.code AS subject_code,b.name AS batch_name
           FROM exam_sessions e JOIN batches b ON b.id=e.batch_id
           LEFT JOIN assessments a ON a.id=COALESCE(e.assessment_id,b.assessment_id)
           LEFT JOIN subjects s ON s.id=a.subject_id
           WHERE e.id=?""",
        (sid,),
    ).fetchone()
    if not ex:
        conn.close(); return jsonify({"ok": False, "error": "Session not found."}), 404
    rows = conn.execute("SELECT id,sender,message,created_at FROM exam_messages WHERE session_id=? ORDER BY id", (sid,)).fetchall()
    conn.execute("UPDATE exam_messages SET read_at=? WHERE session_id=? AND sender='student' AND read_at IS NULL", (iso_now(),sid))
    conn.commit(); conn.close()
    return jsonify({"ok": True, "student": dict(ex), "messages": [dict(r) for r in rows]})


@app.post("/admin/messages/<int:sid>/delete")
@admin_required
def admin_message_delete(sid):
    require_csrf()
    conn = connect()
    ex = conn.execute("SELECT id FROM exam_sessions WHERE id=?", (sid,)).fetchone()
    if not ex:
        conn.close(); return jsonify({"ok": False, "error": "Session not found."}), 404
    count = conn.execute("SELECT COUNT(*) FROM exam_messages WHERE session_id=?", (sid,)).fetchone()[0]
    conn.execute("DELETE FROM exam_messages WHERE session_id=?", (sid,))
    conn.commit(); conn.close()
    return jsonify({"ok": True, "deleted": int(count or 0)})


@app.route("/admin/testing")
@admin_required
def admin_testing():
    conn = connect()
    assessment_options, selected_assessment = _resolve_admin_assessment(conn)
    if not selected_assessment:
        conn.close(); abort(404)
    assessment_id = int(selected_assessment["id"])
    max_score = _assessment_max_score(conn, selected_assessment)
    if selected_assessment["assessment_type"] == "programming_lab":
        lab = conn.execute(
            "SELECT id,title FROM programming_labs WHERE assessment_id=? LIMIT 1",
            (assessment_id,),
        ).fetchone()
        conn.close()
        return render_template(
            "admin_testing_programming.html", assessment_options=assessment_options,
            selected_assessment=selected_assessment, assessment="", max_score=max_score, lab=lab
        )
    batches = conn.execute("SELECT * FROM batches WHERE assessment_id=? ORDER BY slot", (assessment_id,)).fetchall()
    tests = conn.execute(
        """SELECT e.*, b.name AS batch_name FROM exam_sessions e
           JOIN batches b ON b.id=e.batch_id
           WHERE COALESCE(e.is_test,0)=1 AND e.assessment_id=? ORDER BY e.id DESC LIMIT 50""",
        (assessment_id,),
    ).fetchall()
    conn.close()
    return render_template(
        "admin_testing.html", batches=batches, tests=tests, assessment=selected_assessment["assessment_type"],
        selected_assessment=selected_assessment, assessment_options=assessment_options, max_score=max_score
    )


@app.route("/admin/student-preview", methods=["POST"])
@admin_required
def admin_student_preview():
    """Open the real student UI for the selected objective assessment using an instructor-only test identity."""
    require_csrf()
    raw_id = request.form.get("assessment_id", "").strip()
    if not raw_id.isdigit():
        abort(400)
    conn = connect()
    assessment_row = conn.execute("SELECT * FROM assessments WHERE id=? AND deleted_at IS NULL", (int(raw_id),)).fetchone()
    if not assessment_row or assessment_row["assessment_type"] == "programming_lab":
        conn.close(); abort(404)
    batch = conn.execute("SELECT * FROM batches WHERE assessment_id=? ORDER BY slot LIMIT 1", (int(raw_id),)).fetchone()
    if not batch:
        conn.close(); abort(404)
    stamp = datetime.now(APP_TZ).strftime("%Y%m%d%H%M%S%f")
    preview_email = f"custos.instructor.preview.{raw_id}.{stamp}@test.only"
    cur = conn.execute(
        """INSERT INTO exam_sessions(email,first_name,last_name,student_name,program,class_section,batch_id,assessment_id,started_at,ip_address,user_agent,is_test,test_label,untimed,terms_accepted_at)
           VALUES(?,?,?,?,?,?,?,?,?,?,?,1,?,1,?) RETURNING id""",
        (preview_email, "Instructor", "Preview", "Instructor Preview", "ZT", "11", batch["id"], batch["assessment_id"], iso_now(), request.remote_addr, request.headers.get("User-Agent", "")[:500], f"Instructor Preview · {assessment_row['title']}", iso_now()),
    )
    sid = cur.fetchone()[0]
    try:
        assign_questions_to_session(conn, sid, batch["slot"])
    except ValueError as exc:
        conn.execute("DELETE FROM exam_sessions WHERE id=?", (sid,)); conn.commit(); conn.close()
        flash(str(exc), "error")
        return redirect(url_for("admin_questions", assessment_id=raw_id))
    conn.execute(
        "INSERT INTO proctor_events(session_id,event_type,detail,created_at) VALUES (?,?,?,?)",
        (sid, "instructor_student_preview", "Instructor opened the selected assessment through Student View preview.", iso_now()),
    )
    conn.commit(); conn.close()
    session["student_session_id"] = sid
    return redirect(url_for("exam"))


@app.route("/admin/testing/start", methods=["POST"])
@admin_required
def admin_testing_start():
    require_csrf()
    try:
        batch_id = int(request.form.get("batch_id", ""))
    except ValueError:
        abort(400)
    conn = connect()
    batch = conn.execute("SELECT * FROM batches WHERE id=?", (batch_id,)).fetchone()
    if not batch:
        conn.close()
        abort(404)
    stamp = datetime.now(APP_TZ).strftime("%Y%m%d-%H%M%S-%f")
    email = f"instructor.test.{stamp}@{ALLOWED_EMAIL_DOMAIN}"
    label = request.form.get("test_label", "").strip()[:120] or f"Instructor test · {batch['name']}"
    untimed = 1 if request.form.get("untimed") == "1" else 0
    cur = conn.execute(
        """INSERT INTO exam_sessions(email,batch_id,assessment_id,started_at,ip_address,user_agent,is_test,test_label,untimed)
           VALUES (?,?,?,?,?,?,1,?,?) RETURNING id""",
        (email, batch_id, batch["assessment_id"], iso_now(), request.remote_addr, request.headers.get("User-Agent", "")[:500], label, untimed),
    )
    sid = cur.fetchone()[0]
    try:
        assign_questions_to_session(conn, sid, batch["slot"])
    except ValueError as exc:
        conn.execute("DELETE FROM exam_sessions WHERE id=?", (sid,))
        conn.commit()
        conn.close()
        flash(str(exc), "error")
        return redirect(url_for("admin_questions", assessment_id=batch["assessment_id"], batch_slot=batch["slot"]))
    conn.execute(
        "INSERT INTO proctor_events(session_id,event_type,detail,created_at) VALUES (?,?,?,?)",
        (sid, "test_session_started", "Instructor launched an untimed trial" if untimed else "Instructor launched a testing session", iso_now()),
    )
    conn.commit()
    conn.close()
    session["student_session_id"] = sid
    return redirect(url_for("exam"))


@app.route("/admin/testing/resume/<int:sid>", methods=["POST"])
@admin_required
def admin_testing_resume(sid):
    require_csrf()
    conn = connect()
    row = conn.execute(
        """SELECT e.id, e.is_test, e.status, e.assessment_id, b.assessment_type
           FROM exam_sessions e JOIN batches b ON b.id=e.batch_id WHERE e.id=?""",
        (sid,),
    ).fetchone()
    conn.close()
    if not row or not row["is_test"]:
        abort(404)
    if row["status"] == "submitted":
        flash("That trial has already been submitted. Launch a new trial to continue testing.", "error")
        return redirect(url_for("admin_testing", assessment_id=row["assessment_id"]))
    session["student_session_id"] = sid
    return redirect(url_for("exam"))


@app.route("/admin/testing/delete/<int:sid>", methods=["POST"])
@admin_required
def admin_testing_delete(sid):
    require_csrf()
    conn = connect()
    row = conn.execute(
        """SELECT e.is_test, e.assessment_id, b.assessment_type FROM exam_sessions e
           JOIN batches b ON b.id=e.batch_id WHERE e.id=?""",
        (sid,),
    ).fetchone()
    assessment_id = row["assessment_id"] if row else None
    if row and row["is_test"]:
        conn.execute("DELETE FROM exam_sessions WHERE id=?", (sid,))
        conn.commit()
        if session.get("student_session_id") == sid:
            session.pop("student_session_id", None)
        flash("Testing session deleted.", "success")
    conn.close()
    return redirect(url_for("admin_testing", assessment_id=assessment_id) if assessment_id else url_for("admin_testing"))


@app.route("/admin/testing/return", methods=["POST"])
@admin_required
def admin_testing_return():
    require_csrf()
    sid = session.get("student_session_id")
    assessment_id = None
    if sid:
        conn = connect()
        ex = conn.execute(
            """SELECT e.is_test, e.assessment_id, b.assessment_type FROM exam_sessions e
               JOIN batches b ON b.id=e.batch_id WHERE e.id=?""",
            (sid,),
        ).fetchone()
        if ex:
            assessment_id = ex["assessment_id"]
        if ex and ex["is_test"]:
            conn.execute(
                "INSERT INTO proctor_events(session_id,event_type,detail,created_at) VALUES (?,?,?,?)",
                (sid, "test_return_to_admin", "Instructor returned to instructor view", iso_now()),
            )
            conn.commit()
        conn.close()
        session.pop("student_session_id", None)
    return redirect(url_for("admin_testing", assessment_id=assessment_id) if assessment_id else url_for("admin_testing"))


def _live_monitor_payload(conn):
    rows = conn.execute(
        """SELECT e.id,e.email,e.first_name,e.last_name,e.student_name,e.program,e.class_section,e.status,e.started_at,
                  e.flagged_count,e.violation_count,e.security_locked,e.temp_locked_until,e.pending_blackout,e.security_resume_required,e.monitor_done,
                  e.assessment_id, a.title AS assessment_title, b.name AS batch_name,b.assessment_type,
                  COALESCE(msg.unread_messages,0) AS unread_messages,
                  pe.event_type AS last_event,pe.detail AS last_event_detail,pe.created_at AS last_event_at,
                  COALESCE(ans.answered_mcq,0) AS answered_mcq,COALESCE(ans.total_mcq,0) AS total_mcq
           FROM exam_sessions e JOIN batches b ON b.id=e.batch_id
           LEFT JOIN assessments a ON a.id=e.assessment_id
           LEFT JOIN (
               SELECT session_id,COUNT(*) AS unread_messages
               FROM exam_messages
               WHERE sender='student' AND read_at IS NULL
               GROUP BY session_id
           ) msg ON msg.session_id=e.id
           LEFT JOIN (
               SELECT p.session_id,p.event_type,p.detail,p.created_at
               FROM proctor_events p
               JOIN (
                   SELECT session_id,MAX(id) AS max_id
                   FROM proctor_events
                   WHERE event_type<>'behavior_snapshot'
                   GROUP BY session_id
               ) latest ON latest.max_id=p.id
           ) pe ON pe.session_id=e.id
           LEFT JOIN (
               SELECT session_id,
                      SUM(CASE WHEN selected_option IS NOT NULL THEN 1 ELSE 0 END) AS answered_mcq,
                      COUNT(*) AS total_mcq
               FROM session_questions
               GROUP BY session_id
           ) ans ON ans.session_id=e.id
           WHERE e.status='in_progress' AND COALESCE(e.is_test,0)=0 AND COALESCE(e.monitor_done,0)=0
           ORDER BY e.id DESC"""
    ).fetchall()
    rows = instructor_scope.filter_rows(conn, rows, key="assessment_id")
    payload=[]
    for r in rows:
        item=dict(r)
        violations=int(item.get("violation_count") or 0)
        flags=int(item.get("flagged_count") or 0)
        unread=int(item.get("unread_messages") or 0)
        temp_active=False
        temp_until=parse_iso(item.get("temp_locked_until"))
        if temp_until and temp_until > datetime.now(APP_TZ):
            temp_active=True
        risk=0
        if item.get("security_locked"):
            risk += 100
        if item.get("security_resume_required"):
            risk += 70
        if item.get("pending_blackout") or temp_active:
            risk += 45
        risk += min(violations,3)*25
        risk += min(flags,10)*3
        if unread:
            risk += 8
        item["risk_score"]=risk
        if item.get("security_locked"):
            item["attention_level"]="locked"
            item["attention_label"]="Locked · instructor required"
        elif item.get("security_resume_required"):
            item["attention_level"]="high"
            item["attention_label"]="Unlock granted · secure re-entry pending"
        elif risk >= 65:
            item["attention_level"]="high"
            item["attention_label"]="Needs attention"
        elif risk >= 25:
            item["attention_level"]="watch"
            item["attention_label"]="Watch"
        else:
            item["attention_level"]="clear"
            item["attention_label"]="Clear"
        item["answered"] = int(item.get("answered_mcq") or 0)
        item["total"] = int(item.get("total_mcq") or 0)
        payload.append(item)
    payload.sort(key=lambda x: (-x["risk_score"], x.get("last_name") or x.get("student_name") or x.get("email") or "", x.get("first_name") or ""))
    return payload


@app.route("/admin/monitor")
@admin_required
def admin_monitor():
    conn=connect()
    students=_live_monitor_payload(conn)
    assessment_options=instructor_scope.filter_rows(conn, _admin_assessment_options(conn))
    requested_id = str(request.args.get("assessment_id") or session.get("admin_assessment_id") or "").strip()
    selected_assessment = next((row for row in assessment_options if str(row["id"]) == requested_id), None)
    if selected_assessment:
        students = [row for row in students if int(row.get("assessment_id") or 0) == int(selected_assessment["id"])]
        session["admin_assessment_id"] = int(selected_assessment["id"])
    conn.close()
    attention_count=sum(1 for s in students if s["attention_level"] in {"high","locked"})
    locked_count=sum(1 for s in students if s["attention_level"] == "locked")
    return render_template("admin_monitor.html", students=students, assessment=(selected_assessment["assessment_type"] if selected_assessment else session.get("admin_assessment", "posttest")), assessment_options=assessment_options, selected_assessment=selected_assessment, attention_count=attention_count, locked_count=locked_count)


@app.route("/admin/monitor/data")
@admin_required
def admin_monitor_data():
    conn=connect()
    students=_live_monitor_payload(conn)
    requested_id = str(request.args.get("assessment_id") or "").strip()
    if requested_id.isdigit():
        students = [row for row in students if int(row.get("assessment_id") or 0) == int(requested_id)]
    conn.close()
    return jsonify({"ok": True, "students": students, "updated_at": iso_now()})


@app.post("/admin/monitor/session/<int:sid>/action")
@admin_required
def admin_monitor_action(sid):
    require_csrf()
    data = request.get_json(silent=True) or request.form
    action = str(data.get("action", "")).strip().lower()
    conn = connect()
    ex = conn.execute("SELECT * FROM exam_sessions WHERE id=?", (sid,)).fetchone()
    if not ex:
        conn.close()
        return jsonify({"ok": False, "error": "Session not found."}), 404
    if action == "clear_security":
        conn.execute(
            "UPDATE exam_sessions SET security_locked=0,temp_locked_until=NULL,pending_blackout=0,security_resume_required=1 WHERE id=?",
            (sid,),
        )
        detail = f"Instructor granted another chance; violation count retained at {ex['violation_count'] or 0}. Secure-mode re-entry is required before answering can resume."
        event_type = "monitor_security_cleared"
        message = "Unlock granted. The student must re-enter secure mode before continuing."
    elif action == "mark_done":
        conn.execute("UPDATE exam_sessions SET monitor_done=1 WHERE id=?", (sid,))
        detail = "Instructor marked this attempt as done in Live Monitor. Exam status and answers were not changed."
        event_type = "monitor_marked_done"
        message = "Student removed from Live Monitor."
    elif action == "restore":
        conn.execute("UPDATE exam_sessions SET monitor_done=0 WHERE id=?", (sid,))
        detail = "Instructor returned this attempt to Live Monitor."
        event_type = "monitor_restored"
        message = "Student returned to Live Monitor."
    else:
        conn.close()
        return jsonify({"ok": False, "error": "Unknown monitor action."}), 400
    conn.execute(
        "INSERT INTO proctor_events(session_id,event_type,detail,created_at) VALUES (?,?,?,?)",
        (sid,event_type,detail,iso_now()),
    )
    conn.commit(); conn.close()
    return jsonify({"ok": True, "message": message})



def _admin_assessment_options(conn):
    return conn.execute(
        """SELECT a.*, s.code AS subject_code, s.name AS subject_name,
                  (SELECT COUNT(*) FROM questions q WHERE q.assessment_id=a.id AND COALESCE(q.active,1)=1) AS question_count
           FROM assessments a JOIN subjects s ON s.id=a.subject_id
           WHERE a.deleted_at IS NULL
           ORDER BY a.active DESC,
                    CASE WHEN a.assessment_type='programming_lab' THEN 1 ELSE 0 END,
                    a.created_at DESC, a.id DESC"""
    ).fetchall()


def _resolve_admin_assessment(conn):
    options = instructor_scope.filter_rows(conn, _admin_assessment_options(conn))
    if not options:
        return options, None
    raw_id = request.args.get("assessment_id", "").strip()
    selected = None
    if raw_id.isdigit():
        selected = next((row for row in options if int(row["id"]) == int(raw_id)), None)
    if selected is None:
        legacy = request.args.get("assessment", "").strip().lower()
        if legacy:
            selected = next((row for row in options if row["assessment_type"] == legacy), None)
    if selected is None:
        remembered = session.get("admin_assessment_id")
        if remembered:
            selected = next((row for row in options if int(row["id"]) == int(remembered)), None)
    if selected is None:
        selected = next((row for row in options if row["assessment_type"] != "programming_lab"), options[0])
    session["admin_assessment_id"] = int(selected["id"])
    if selected["assessment_type"] in {"midterm", "posttest"}:
        session["admin_assessment"] = selected["assessment_type"]
    return options, selected


def _assessment_max_score(conn, assessment_row):
    if not assessment_row:
        return 0
    atype = assessment_row["assessment_type"]
    if atype == "midterm":
        return 55
    if atype == "posttest":
        return 50
    if atype == "programming_lab":
        row = conn.execute(
            """SELECT COALESCE(SUM(t.points),0) AS total
               FROM programming_tasks t JOIN programming_labs pl ON pl.id=t.lab_id
               WHERE pl.assessment_id=? AND COALESCE(t.active,1)=1""",
            (assessment_row["id"],),
        ).fetchone()
        return float(row["total"] or 0)
    point_rows = conn.execute(
        "SELECT COALESCE(points,1) AS points FROM questions WHERE assessment_id=? AND COALESCE(active,1)=1 ORDER BY COALESCE(points,1) DESC",
        (assessment_row["id"],),
    ).fetchall()
    points = [float(row["points"] or 1) for row in point_rows]
    limit = int(assessment_row["question_limit"] or 0) if "question_limit" in assessment_row.keys() else 0
    if limit > 0:
        points = points[:limit]
    total = sum(points)
    return int(total) if float(total).is_integer() else round(total, 2)

@app.route("/admin/questions")
@admin_required
def admin_questions():
    conn = connect()
    assessment_options, selected_assessment = _resolve_admin_assessment(conn)
    if not selected_assessment:
        conn.close(); abort(404)

    assessment_id = int(selected_assessment["id"])
    atype = selected_assessment["assessment_type"]
    search = re.sub(r"\s+", " ", request.args.get("q", "").strip())[:120]
    status_filter = request.args.get("status", "all").strip().lower()
    if status_filter not in {"all", "active", "disabled"}:
        status_filter = "all"

    if atype == "programming_lab":
        tasks = conn.execute(
            """SELECT t.* FROM programming_tasks t JOIN programming_labs pl ON pl.id=t.lab_id
               WHERE pl.assessment_id=? ORDER BY t.position,t.id""",
            (assessment_id,),
        ).fetchall()
        conn.close()
        return render_template(
            "admin_questions_lab.html", assessment_options=assessment_options,
            selected_assessment=selected_assessment, tasks=tasks, assessment=""
        )

    if atype == "custom" or atype == "dryrun":
        params = [assessment_id]
        where = ["q.assessment_id=?"]
        if search:
            term = f"%{search}%"
            where.append("(q.topic LIKE ? OR q.prompt LIKE ? OR q.code LIKE ?)")
            params.extend([term, term, term])
        if status_filter == "active":
            where.append("COALESCE(q.active,1)=1")
        elif status_filter == "disabled":
            where.append("COALESCE(q.active,1)=0")
        questions = conn.execute(
            f"""SELECT q.* FROM questions q WHERE {' AND '.join(where)}
                 ORDER BY COALESCE(q.position,q.id),q.id""",
            params,
        ).fetchall()
        counts = conn.execute(
            """SELECT COUNT(*) AS total,
                      SUM(CASE WHEN COALESCE(active,1)=1 THEN 1 ELSE 0 END) AS active_count,
                      COALESCE(SUM(CASE WHEN COALESCE(active,1)=1 THEN COALESCE(points,1) ELSE 0 END),0) AS active_points
               FROM questions WHERE assessment_id=?""",
            (assessment_id,),
        ).fetchone()
        conn.close()
        return render_template(
            "admin_questions_custom.html", assessment_options=assessment_options,
            selected_assessment=selected_assessment, questions=questions, counts=counts,
            search=search, status_filter=status_filter, assessment=""
        )

    assessment = atype if atype in {"midterm", "posttest"} else "posttest"
    batches = conn.execute(
        "SELECT slot,name FROM batches WHERE assessment_id=? ORDER BY slot", (assessment_id,)
    ).fetchall()
    if not batches:
        conn.close(); abort(404)
    requested_slot = request.args.get("batch_slot", "").strip()
    selected_slot = int(requested_slot) if requested_slot.isdigit() else batches[0]["slot"]
    valid_slots = {b["slot"] for b in batches}
    if selected_slot not in valid_slots:
        selected_slot = batches[0]["slot"]

    source_filter = request.args.get("source", "all").strip().lower()
    if source_filter not in {"all", "builtin", "instructor"}:
        source_filter = "all"
    params = [selected_slot]
    where = ["q.batch_slot=?"]
    if search:
        where.append("(q.topic LIKE ? OR q.prompt LIKE ? OR q.code LIKE ?)")
        term = f"%{search}%"; params.extend([term,term,term])
    if status_filter == "active": where.append("COALESCE(q.active,1)=1")
    elif status_filter == "disabled": where.append("COALESCE(q.active,1)=0")
    if source_filter == "instructor": where.append("q.created_by='instructor'")
    elif source_filter == "builtin": where.append("q.created_by!='instructor'")
    rows = conn.execute(
        f"""SELECT q.*, b.name AS batch_name FROM questions q JOIN batches b ON b.slot=q.batch_slot
             WHERE {' AND '.join(where)} ORDER BY q.part,q.id""", params
    ).fetchall()
    questions=[dict(r) for r in rows]
    part1_questions=[q for q in questions if q["part"]==1]
    part2_questions=[q for q in questions if q["part"]==2]
    counts=conn.execute(
        """SELECT part,COUNT(*) AS total,SUM(CASE WHEN COALESCE(active,1)=1 THEN 1 ELSE 0 END) AS active_count
           FROM questions WHERE batch_slot=? GROUP BY part ORDER BY part""", (selected_slot,)
    ).fetchall()
    bonus=conn.execute("SELECT * FROM bonus_questions WHERE assessment_type='midterm' ORDER BY position").fetchall() if assessment=="midterm" else []
    conn.close()
    return render_template(
        "admin_questions.html", assessment=assessment, assessment_options=assessment_options,
        selected_assessment=selected_assessment, batches=batches, selected_slot=selected_slot,
        questions=questions, part1_questions=part1_questions, part2_questions=part2_questions,
        counts={r["part"]:dict(r) for r in counts}, bonus=bonus, search=search,
        status_filter=status_filter, source_filter=source_filter
    )


def _question_form_values(form):
    try:
        batch_slot = int(form.get("batch_slot", ""))
        part = int(form.get("part", ""))
    except (TypeError, ValueError):
        raise ValueError("Choose a valid bank and part.")
    if part not in {1, 2}:
        raise ValueError("Part must be Part I or Part II.")
    topic = re.sub(r"\s+", " ", form.get("topic", "").strip())[:120]
    prompt = form.get("prompt", "").strip()[:1200]
    code = form.get("code", "").rstrip()[:5000]
    options = {letter: form.get(f"option_{letter.lower()}", "").strip()[:1000] for letter in "ABCD"}
    correct = form.get("correct_option", "").strip().upper()
    explanation = form.get("explanation", "").strip()[:2000]
    if not topic or not prompt or any(not v for v in options.values()):
        raise ValueError("Topic, prompt, and all four answer choices are required.")
    if len(set(options.values())) != 4:
        raise ValueError("All four answer choices must be different.")
    if correct not in {"A", "B", "C", "D"}:
        raise ValueError("Choose the correct answer key.")
    return batch_slot, part, topic, prompt, code, options, correct, explanation


@app.route("/admin/questions/add", methods=["POST"])
@admin_required
def admin_question_add():
    require_csrf()
    assessment = request.form.get("assessment", "posttest").strip().lower()
    if assessment not in {"midterm", "posttest"}:
        assessment = "posttest"
    try:
        batch_slot, part, topic, prompt, code, options, correct, explanation = _question_form_values(request.form)
        conn = connect()
        batch = conn.execute("SELECT assessment_type,subject_id,assessment_id FROM batches WHERE slot=?", (batch_slot,)).fetchone()
        if not batch or batch["assessment_type"] != assessment:
            conn.close()
            raise ValueError("The selected bank does not belong to this assessment.")
        conn.execute(
            """INSERT INTO questions(part,batch_slot,topic,prompt,code,option_a,option_b,option_c,option_d,correct_option,explanation,active,created_by,subject_id,assessment_id)
               VALUES(?,?,?,?,?,?,?,?,?,?,?,1,'instructor',?,?)""",
            (part, batch_slot, topic, prompt, code, options["A"], options["B"], options["C"], options["D"], correct, explanation, batch["subject_id"], batch["assessment_id"]),
        )
        conn.commit()
        conn.close()
        flash("Question added to the active pool.", "success")
        return redirect(url_for("admin_questions", assessment=assessment, batch_slot=batch_slot, part=part))
    except ValueError as exc:
        flash(str(exc), "error")
        return redirect(url_for("admin_questions", assessment=assessment))


@app.route("/admin/questions/<int:qid>/edit", methods=["GET", "POST"])
@admin_required
def admin_question_edit(qid):
    conn = connect()
    q = conn.execute(
        """SELECT q.*, b.assessment_type, b.name AS batch_name FROM questions q
           JOIN batches b ON b.slot=q.batch_slot WHERE q.id=?""", (qid,)
    ).fetchone()
    if not q:
        conn.close()
        abort(404)
    assessment = q["assessment_type"]
    if request.method == "POST":
        require_csrf()
        try:
            batch_slot, part, topic, prompt, code, options, correct, explanation = _question_form_values(request.form)
            batch = conn.execute("SELECT assessment_type FROM batches WHERE slot=?", (batch_slot,)).fetchone()
            if not batch or batch["assessment_type"] != assessment:
                raise ValueError("The selected bank does not belong to this assessment.")
            conn.execute(
                """UPDATE questions SET part=?,batch_slot=?,topic=?,prompt=?,code=?,option_a=?,option_b=?,option_c=?,option_d=?,correct_option=?,explanation=?,subject_id=?,assessment_id=? WHERE id=?""",
                (part, batch_slot, topic, prompt, code, options["A"], options["B"], options["C"], options["D"], correct, explanation, batch["subject_id"], batch["assessment_id"], qid),
            )
            conn.commit()
            flash("Question updated.", "success")
            conn.close()
            return redirect(url_for("admin_questions", assessment=assessment, batch_slot=batch_slot, part=part))
        except ValueError as exc:
            flash(str(exc), "error")
            q = conn.execute(
                """SELECT q.*, b.assessment_type, b.name AS batch_name FROM questions q
                   JOIN batches b ON b.slot=q.batch_slot WHERE q.id=?""", (qid,)
            ).fetchone()
    batches = conn.execute("SELECT slot,name FROM batches WHERE assessment_type=? ORDER BY slot", (assessment,)).fetchall()
    conn.close()
    return render_template("admin_question_edit.html", q=q, batches=batches, assessment=assessment)


@app.route("/admin/questions/<int:qid>/toggle", methods=["POST"])
@admin_required
def admin_question_toggle(qid):
    require_csrf()
    conn = connect()
    q = conn.execute(
        """SELECT q.*, b.assessment_type FROM questions q JOIN batches b ON b.slot=q.batch_slot WHERE q.id=?""",
        (qid,),
    ).fetchone()
    if not q:
        conn.close()
        abort(404)
    new_active = 0 if q["active"] else 1
    if q["active"] and new_active == 0:
        active_count = conn.execute(
            "SELECT COUNT(*) FROM questions WHERE batch_slot=? AND part=? AND active=1", (q["batch_slot"], q["part"])
        ).fetchone()[0]
        required = 40 if q["part"] == 1 else 20
        if active_count <= required:
            conn.close()
            flash(f"Cannot disable this item: the bank must keep at least {required} active Part {q['part']} questions.", "error")
            return redirect(url_for("admin_questions", assessment=q["assessment_type"], batch_slot=q["batch_slot"], part=q["part"]))
    conn.execute("UPDATE questions SET active=? WHERE id=?", (new_active, qid))
    conn.commit()
    conn.close()
    flash("Question enabled." if new_active else "Question disabled from future sessions.", "success")
    return redirect(url_for("admin_questions", assessment=q["assessment_type"], batch_slot=q["batch_slot"], part=q["part"]))


@app.route("/admin/bonus/<int:qid>/edit", methods=["POST"])
@admin_required
def admin_bonus_edit(qid):
    require_csrf()
    prompt = request.form.get("prompt", "").strip()[:1200]
    accepted = request.form.get("accepted_answer", "").strip()[:300]
    topic = re.sub(r"\s+", " ", request.form.get("topic", "General Knowledge").strip())[:120]
    if not prompt or not accepted or not any(normalize_short_answer(x) for x in accepted.split("|")):
        flash("Bonus prompt and a short accepted answer are required.", "error")
        return redirect(url_for("admin_questions", assessment="midterm"))
    conn = connect()
    row = conn.execute("SELECT id FROM bonus_questions WHERE id=? AND assessment_type='midterm'", (qid,)).fetchone()
    if not row:
        conn.close()
        abort(404)
    conn.execute("UPDATE bonus_questions SET prompt=?,accepted_answer=?,topic=? WHERE id=?", (prompt, accepted, topic, qid))
    conn.commit()
    conn.close()
    flash("Bonus question updated. Answer checking ignores case, spaces, dashes, and punctuation.", "success")
    return redirect(url_for("admin_questions", assessment="midterm"))


@app.route("/admin/analysis")
@admin_required
def admin_analysis():
    try:
        return _admin_analysis_impl()
    except Exception as exc:
        # Only database-driver errors should trigger an automatic schema repair.
        # Template/programming errors cannot be fixed by a migration and retrying
        # them adds needless latency on serverless deployments.
        exc_module = type(exc).__module__.lower()
        is_database_error = exc_module.startswith("sqlite3") or exc_module.startswith("psycopg")
        if not is_database_error:
            app.logger.exception("Item Analysis failed")
            return render_template(
                "admin_analysis_error.html",
                error_type=type(exc).__name__,
                assessment=session.get("admin_assessment", "posttest"),
            ), 500

        app.logger.exception("Item Analysis database error; attempting one schema repair retry")
        try:
            ensure_db_initialized(
                admin_username=os.getenv("ADMIN_USERNAME", "admin"),
                admin_password=_admin_password,
                force=True,
            )
            return _admin_analysis_impl()
        except Exception:
            app.logger.exception("Item Analysis retry failed")
            # Do not expose database credentials, SQL, or a traceback in the UI.
            return render_template(
                "admin_analysis_error.html",
                error_type=type(exc).__name__,
                assessment=session.get("admin_assessment", "posttest"),
            ), 500


def _admin_analysis_impl():
    conn = connect()
    assessment_options, selected_assessment = _resolve_admin_assessment(conn)
    if not selected_assessment:
        conn.close(); abort(404)
    assessment_id = int(selected_assessment["id"])
    atype = selected_assessment["assessment_type"]
    max_score = _assessment_max_score(conn, selected_assessment)

    if atype == "programming_lab":
        analysis = build_programming_analysis(conn, assessment_id=assessment_id, max_score=max_score)
        conn.close()
        return render_template(
            "admin_analysis_programming.html",
            assessment_options=assessment_options,
            selected_assessment=selected_assessment,
            analysis=analysis,
            max_score=max_score,
            assessment="",
        )

    batch_slot = request.args.get("batch_slot", "").strip()
    part = request.args.get("part", "").strip()
    topic = request.args.get("topic", "").strip()
    analysis = build_item_analysis(
        conn,
        batch_slot=int(batch_slot) if batch_slot.isdigit() else None,
        part=int(part) if part.isdigit() else None,
        topic=topic or None,
        assessment_id=assessment_id,
        assessment=atype,
        max_score=max_score,
    )
    # Keep filter metadata compatible with both current and older databases.
    # Legacy attempts are identified by assessment type when their historical
    # assessment_id link has not yet been populated.
    batches = conn.execute(
        """SELECT slot,name FROM batches
           WHERE assessment_id=? OR (assessment_id IS NULL AND assessment_type=?)
           ORDER BY slot""",
        (assessment_id, atype),
    ).fetchall()
    topics = [dict(r) for r in conn.execute(
        """SELECT DISTINCT q.topic
           FROM questions q
           LEFT JOIN batches b ON b.slot=q.batch_slot
           WHERE (q.assessment_id=? OR (q.assessment_id IS NULL AND b.assessment_type=?))
             AND q.topic IS NOT NULL AND q.topic<>''
           ORDER BY q.topic""",
        (assessment_id, atype),
    ).fetchall()]
    if atype == "midterm":
        known_topics = {str(r.get("topic") or "") for r in topics}
        for r in conn.execute(
            "SELECT DISTINCT topic FROM bonus_questions WHERE assessment_type='midterm' AND topic IS NOT NULL AND topic<>'' ORDER BY topic"
        ).fetchall():
            value = str(r[0] or "")
            if value and value not in known_topics:
                topics.append({"topic": value})
                known_topics.add(value)
        topics.sort(key=lambda row: str(row.get("topic") or "").casefold())
    parts = [int(r[0]) for r in conn.execute(
        """SELECT DISTINCT q.part
           FROM questions q
           LEFT JOIN batches b ON b.slot=q.batch_slot
           WHERE q.assessment_id=? OR (q.assessment_id IS NULL AND b.assessment_type=?)
           ORDER BY q.part""",
        (assessment_id, atype),
    ).fetchall()]
    if atype == "midterm" and 3 not in parts:
        parts.append(3)
    conn.close()
    return render_template(
        "admin_analysis.html",
        analysis=analysis,
        batches=batches,
        topics=topics,
        parts=parts,
        selected_batch=batch_slot,
        selected_part=part,
        selected_topic=topic,
        assessment=atype,
        assessment_options=assessment_options,
        selected_assessment=selected_assessment,
        max_score=max_score,
        is_legacy=atype in {"midterm", "posttest", "dryrun"},
    )


@app.route("/admin/analysis/item/<int:qid>")
@admin_required
def admin_analysis_item(qid):
    conn = connect()
    question = conn.execute("SELECT * FROM questions WHERE id=?", (qid,)).fetchone()
    if not question:
        conn.close(); abort(404)
    selected_assessment = conn.execute(
        """SELECT a.*,s.code AS subject_code,s.name AS subject_name
           FROM assessments a JOIN subjects s ON s.id=a.subject_id
           WHERE a.id=? AND a.deleted_at IS NULL""",
        (question["assessment_id"],),
    ).fetchone()
    if not selected_assessment:
        conn.close(); abort(404)
    max_score = _assessment_max_score(conn, selected_assessment)
    analysis = build_item_analysis(
        conn,
        assessment_id=selected_assessment["id"],
        assessment=selected_assessment["assessment_type"],
        max_score=max_score,
    )
    item = next((i for i in analysis["items"] if i.get("question_id") == qid), None)
    conn.close()
    return render_template(
        "admin_analysis_item.html",
        item=item,
        question=question,
        assessment=selected_assessment["assessment_type"],
        selected_assessment=selected_assessment,
        is_legacy=selected_assessment["assessment_type"] in {"midterm", "posttest", "dryrun"},
    )


@app.route("/admin/analysis.csv")
@admin_required
def export_item_analysis():
    conn = connect()
    assessment_options, selected_assessment = _resolve_admin_assessment(conn)
    if not selected_assessment:
        conn.close(); abort(404)
    assessment_id = int(selected_assessment["id"])
    atype = selected_assessment["assessment_type"]
    max_score = _assessment_max_score(conn, selected_assessment)
    sio = io.StringIO()
    writer = csv.writer(sio)

    if atype == "programming_lab":
        analysis = build_programming_analysis(conn, assessment_id=assessment_id, max_score=max_score)
        writer.writerow([
            "assessment", "task", "title", "points", "students", "mean_score",
            "mean_attainment_pct", "avg_runs", "full_score_pct", "zero_score_pct",
            "corrected_task_total_r", "high_low_27", "flags",
        ])
        for item in analysis["tasks"]:
            writer.writerow([
                selected_assessment["title"], item["item_id"], item["title"], item["points"], item["n"],
                item["avg_score"], item["attainment_pct"], item["avg_runs"], item["full_score_pct"],
                item["zero_score_pct"], item["r_pb_display"], item["high_low_display"], " | ".join(item["flags"]),
            ])
    else:
        analysis = build_item_analysis(
            conn,
            assessment_id=assessment_id,
            assessment=atype,
            max_score=max_score,
        )
        writer.writerow([
            "assessment", "item", "type", "set", "part", "topic", "responses", "correct",
            "percent_correct", "omitted", "omission_pct", "corrected_point_biserial_r",
            "high_low_27", "A", "B", "C", "D", "correct_key", "functioning_distractors", "flags",
        ])
        for item in analysis["items"]:
            counts = item.get("option_counts") or {}
            writer.writerow([
                selected_assessment["title"], item["item_id"], item["item_type"], item.get("batch_slot") or "all",
                item["part"], item["topic"], item["n"], item["correct"], item["p_pct"], item["omitted"],
                item["omission_pct"], item["r_pb_display"], item["high_low_display"], counts.get("A", ""),
                counts.get("B", ""), counts.get("C", ""), counts.get("D", ""), item.get("correct_option") or "",
                item.get("functioning_distractors") if item.get("functioning_distractors") is not None else "",
                " | ".join(item["flags"]),
            ])
    conn.close()
    safe = re.sub(r"[^A-Za-z0-9_-]+", "-", selected_assessment["title"]).strip("-") or "assessment"
    return Response(
        sio.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": f"attachment; filename={safe}-item-analysis.csv"},
    )


@app.route("/admin/export/session-keys.csv")
@admin_required
def export_session_keys():
    assessment = request.args.get("assessment", "midterm").strip().lower()
    if assessment not in {"midterm", "posttest"}:
        assessment = "midterm"
    conn = connect()
    rows = conn.execute(
        """SELECT slot, section, batch_label, name, access_code, active, open_at, close_at
           FROM batches WHERE assessment_type=? ORDER BY slot""",
        (assessment,),
    ).fetchall()
    conn.close()
    sio = io.StringIO()
    writer = csv.writer(sio)
    writer.writerow(["set", "section", "batch", "name", "session_key", "active", "opens", "closes"])
    for r in rows:
        set_no = r["slot"] - 8 if assessment == "posttest" else r["slot"]
        writer.writerow([set_no, r["section"], r["batch_label"], r["name"], r["access_code"], "yes" if r["active"] else "no", r["open_at"] or "", r["close_at"] or ""])
    return Response(
        sio.getvalue(), mimetype="text/csv",
        headers={"Content-Disposition": f"attachment; filename=csdc101_{assessment}_session_keys.csv"},
    )


PRIVATE_BANK_COLUMNS = {
    "part", "batch_slot", "topic", "prompt", "code",
    "option_a", "option_b", "option_c", "option_d",
    "correct_option", "explanation",
}


def _parse_private_bank_upload(file_storage):
    if not file_storage or not file_storage.filename:
        raise ValueError("Choose a private question-bank CSV file.")
    try:
        raw = file_storage.read().decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("The CSV must be UTF-8 encoded.") from exc
    reader = csv.DictReader(io.StringIO(raw))
    headers = set(reader.fieldnames or [])
    missing = PRIVATE_BANK_COLUMNS - headers
    if missing:
        raise ValueError("Missing CSV columns: " + ", ".join(sorted(missing)))

    rows = []
    for line_no, row in enumerate(reader, start=2):
        try:
            part = int((row.get("part") or "").strip())
            slot = int((row.get("batch_slot") or "").strip())
        except ValueError as exc:
            raise ValueError(f"Line {line_no}: part and batch_slot must be integers.") from exc
        if part not in {1, 2}:
            raise ValueError(f"Line {line_no}: part must be 1 or 2.")
        if slot not in range(1, 17):
            raise ValueError(f"Line {line_no}: batch_slot must be from 1 through 16.")
        correct = (row.get("correct_option") or "").strip().upper()
        if correct not in {"A", "B", "C", "D"}:
            raise ValueError(f"Line {line_no}: correct_option must be A, B, C, or D.")
        values = {key: (row.get(key) or "").strip() for key in PRIVATE_BANK_COLUMNS}
        if not values["topic"] or not values["prompt"]:
            raise ValueError(f"Line {line_no}: topic and prompt are required.")
        if any(not values[f"option_{letter}"] for letter in "abcd"):
            raise ValueError(f"Line {line_no}: all four answer choices are required.")
        if len({values[f"option_{letter}"] for letter in "abcd"}) != 4:
            raise ValueError(f"Line {line_no}: all four answer choices must be different.")
        values["part"] = part
        values["batch_slot"] = slot
        values["correct_option"] = correct
        rows.append(values)
    if not rows:
        raise ValueError("The CSV contains no question rows.")
    return rows


@app.route("/admin/questions/import-private", methods=["POST"])
@admin_required
def import_private_questions():
    require_csrf()
    assessment = request.form.get("assessment", "posttest")
    if request.form.get("confirm_private_import") != "yes":
        flash("Confirm that this is your private question-bank file before importing.", "error")
        return redirect(url_for("admin_questions", assessment=assessment))
    try:
        rows = _parse_private_bank_upload(request.files.get("bank_file"))
        affected_groups = {"midterm" if r["batch_slot"] <= 8 else "posttest" for r in rows}

        included_slots = sorted({r["batch_slot"] for r in rows})
        for slot in included_slots:
            p1 = sum(1 for r in rows if r["batch_slot"] == slot and r["part"] == 1)
            p2 = sum(1 for r in rows if r["batch_slot"] == slot and r["part"] == 2)
            if p1 < 40 or p2 < 20:
                raise ValueError(
                    f"Bank slot {slot} has only {p1} Part I and {p2} Part II items; at least 40/20 are required."
                )

        conn = connect()
        try:
            for group in affected_groups:
                real = conn.execute(
                    """SELECT COUNT(*) FROM exam_sessions e
                       JOIN batches b ON b.id=e.batch_id
                       WHERE COALESCE(e.is_test,0)=0 AND b.assessment_type=?""",
                    (group,),
                ).fetchone()[0]
                if real:
                    raise ValueError(f"Cannot replace the {group} bank after real student attempts exist.")

            conn.execute("DELETE FROM exam_sessions WHERE COALESCE(is_test,0)=1")
            for group in affected_groups:
                lo, hi = (1, 8) if group == "midterm" else (9, 16)
                conn.execute("DELETE FROM questions WHERE batch_slot BETWEEN ? AND ?", (lo, hi))

            conn.executemany(
                """INSERT INTO questions(
                       part,batch_slot,topic,prompt,code,option_a,option_b,option_c,option_d,
                       correct_option,explanation,active,created_by
                   ) VALUES(?,?,?,?,?,?,?,?,?,?,?,1,'private-import')""",
                [
                    (
                        r["part"], r["batch_slot"], r["topic"], r["prompt"], r["code"],
                        r["option_a"], r["option_b"], r["option_c"], r["option_d"],
                        r["correct_option"], r["explanation"],
                    )
                    for r in rows
                ],
            )

            if "posttest" in affected_groups:
                conn.execute("DELETE FROM questions WHERE batch_slot=17")
                conn.execute(
                    """INSERT INTO questions(
                           part,batch_slot,topic,prompt,code,option_a,option_b,option_c,option_d,
                           correct_option,explanation,active,created_by
                       )
                       SELECT part,17,topic,prompt,code,option_a,option_b,option_c,option_d,
                              correct_option,explanation,active,'dryrun-runtime-copy'
                       FROM questions WHERE batch_slot=9 ORDER BY id"""
                )
            conn.execute(
                """UPDATE questions SET subject_id=(SELECT b.subject_id FROM batches b WHERE b.slot=questions.batch_slot),
                       assessment_id=(SELECT b.assessment_id FROM batches b WHERE b.slot=questions.batch_slot)
                   WHERE subject_id IS NULL OR assessment_id IS NULL"""
            )
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            conn.close()
        flash(
            f"Private bank imported: {len(rows)} questions. The uploaded CSV was processed in memory and was not saved by Custos.",
            "success",
        )
    except ValueError as exc:
        flash(str(exc), "error")
    return redirect(url_for("admin_questions", assessment=assessment))


@app.route("/admin/export/questions.csv")
@admin_required
def export_questions():
    assessment = request.args.get("assessment", "midterm").strip().lower()
    if assessment not in {"midterm", "posttest"}:
        assessment = "midterm"
    conn = connect()
    rows = conn.execute(
        """SELECT q.* FROM questions q JOIN batches b ON b.slot=q.batch_slot
           WHERE b.assessment_type=? ORDER BY q.batch_slot, q.part, q.id""",
        (assessment,),
    ).fetchall()
    conn.close()
    sio = io.StringIO()
    writer = csv.writer(sio)
    writer.writerow(["id", "part", "batch_slot", "topic", "prompt", "code", "option_a", "option_b", "option_c", "option_d", "correct_option", "active", "created_by"])
    for r in rows:
        writer.writerow([r[k] for k in ["id", "part", "batch_slot", "topic", "prompt", "code", "option_a", "option_b", "option_c", "option_d", "correct_option", "active", "created_by"]])
    return Response(sio.getvalue(), mimetype="text/csv", headers={"Content-Disposition": f"attachment; filename=csdc101_{assessment}_question_bank.csv"})


# Custos 1.0 · Goliathus: subjects, multi-assessment workspace, and secure C++ Programming Lab.
from workspace import register as register_nextgen
register_nextgen(app)
from google_integration import register as register_google
register_google(app)
import instructor_scope
instructor_scope.register(app)


if __name__ == "__main__":
    host = os.getenv("HOST", "0.0.0.0")
    port = int(os.getenv("PORT", "5000"))
    if os.getenv("FLASK_DEBUG", "0") == "1":
        app.run(host=host, port=port, debug=True)
    else:
        # Waitress is production-grade on Windows and Linux and makes this same
        # build suitable for a classroom LAN server without requiring Gunicorn.
        from waitress import serve
        serve(app, host=host, port=port, threads=int(os.getenv("WAITRESS_THREADS", "8")))
